# -*- coding: utf-8 -*-
"""Carga de productos de la tienda scout desde Excel: Producto · Existencia · Precio de venta.

Si el producto ya existe (mismo nombre) se actualiza su precio y su existencia queda igual a la del archivo
(con un movimiento de inventario inicial o de ajuste), así se puede volver a subir el archivo sin duplicar nada.
"""
import base64
import io
import logging
import re
import unicodedata

from odoo import _, api, fields, models
from odoo.exceptions import UserError
from odoo.tools import float_compare

_logger = logging.getLogger(__name__)

HEADERS = {
    'name': ('producto', 'nombre', 'descripcion', 'nombre del producto'),
    'qty': ('existencia', 'cantidad', 'existencias', 'stock', 'inventario'),
    'price': ('precio de venta', 'precio', 'precio venta', 'precio unitario'),
}


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', text).strip().rstrip(' *')


def _number(value, label):
    if value in (None, ''):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace('$', '').replace(',', '').strip()
    try:
        return float(text)
    except ValueError:
        raise ValueError(_('%(label)s "%(value)s" no es un número.', label=label, value=value))


class ScoutProductImport(models.TransientModel):
    _name = 'scout.product.import'
    _description = 'Carga de productos de la tienda scout'

    file = fields.Binary(string='Archivo Excel (.xlsx)')
    file_name = fields.Char()
    template_file = fields.Binary(string='Plantilla', readonly=True, compute='_compute_template_file')
    template_name = fields.Char(default='plantilla_productos_tienda.xlsx')
    inventory_date = fields.Date(string='Fecha del inventario', required=True, default=fields.Date.context_today,
                                 help='Fecha con la que se registra la existencia (inventario inicial o ajuste).')
    state = fields.Selection([('draft', 'Cargar'), ('checked', 'Revisado'), ('done', 'Importado')], default='draft')
    result = fields.Text(string='Resultado', readonly=True)
    created_count = fields.Integer(readonly=True)
    updated_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)

    # ------------------------------------------------------------------ plantilla
    def _compute_template_file(self):
        content = self._build_template()
        for wizard in self:
            wizard.template_file = base64.b64encode(content)

    @api.model
    def _build_template(self):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = 'Productos'
        for col, (label, width) in enumerate((('Producto *', 55), ('Existencia', 14), ('Precio de venta *', 16)), start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='0F3049')
            ws.column_dimensions[cell.column_letter].width = width
        ws.append(['Pañoleta oficial del grupo talla G', 7, 250])
        ws.append(['Kit de promesa (deja Existencia vacía si no controla piezas)', None, 360])
        ws.freeze_panes = 'A2'
        info = wb.create_sheet('Instrucciones')
        for row, text in enumerate([
            'PRODUCTOS DE LA TIENDA SCOUT',
            'Una fila por producto: Producto, Existencia y Precio de venta. Las filas 2 y 3 son ejemplos: bórralas.',
            'Si el producto ya existe (mismo nombre) se actualiza: su precio cambia y su existencia queda igual a la del archivo.',
            'Existencia vacía = el producto no controla existencias (cuotas, alimentos preparados, kits).',
            'Los productos quedan sin IVA y como artículos de la tienda scout.',
        ], start=1):
            info.cell(row=row, column=1, value=text).font = Font(bold=(row == 1))
        info.column_dimensions['A'].width = 120
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    # ------------------------------------------------------------------ lectura
    def _read_rows(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('Selecciona el archivo de Excel.'))
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(base64.b64decode(self.file)), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001
            raise UserError(_('No se pudo leer el archivo. Usa un archivo .xlsx. (%s)', exc)) from exc
        ws = wb['Productos'] if 'Productos' in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        # Busca la fila de encabezados (puede haber títulos arriba, como en el inventario de la tiendita)
        for start, row in enumerate(rows[:15]):
            header = [_norm(h) for h in row]
            index = {}
            for key, names in HEADERS.items():
                for i, h in enumerate(header):
                    if h in names or (key == 'price' and h.startswith('precio')):
                        index.setdefault(key, i)
            if 'name' in index and ('price' in index or 'qty' in index):
                break
        else:
            raise UserError(_('No encontré las columnas "Producto", "Existencia" y "Precio de venta".'))
        items = []
        for number, row in enumerate(rows[start + 1:], start=start + 2):
            def get(key):
                i = index.get(key)
                return row[i] if i is not None and i < len(row) else None
            name = re.sub(r'\s+', ' ', str(get('name') or '')).strip()
            if not name or _norm(name).startswith('total'):
                continue
            items.append((number, {'name': name, 'qty': get('qty'), 'price': get('price')}))
        return items

    # ------------------------------------------------------------------ importación
    def _import_row(self, item):
        qty = _number(item['qty'], _('Existencia'))
        price = _number(item['price'], _('Precio de venta'))
        if qty is not None and qty < 0:
            raise ValueError(_('La existencia no puede ser negativa.'))
        if price is not None and price < 0:
            raise ValueError(_('El precio no puede ser negativo.'))
        Template = self.env['product.template'].with_context(active_test=False)
        product = Template.search([('name', '=ilike', item['name']), ('scout_store_ok', '=', True)], limit=1) \
            or Template.search([('name', '=ilike', item['name'])], limit=1)
        created = not product
        if created:
            product = Template.create({
                'name': item['name'], 'type': 'consu', 'sale_ok': True, 'purchase_ok': True,
                'scout_store_ok': True, 'scout_sale_type': 'store', 'scout_track_qty': qty is not None,
                'list_price': price or 0.0, 'taxes_id': [fields.Command.clear()],
                'supplier_taxes_id': [fields.Command.clear()],
            })
        else:
            vals = {'scout_store_ok': True, 'active': True}
            if price is not None:
                vals['list_price'] = price
            if qty is not None:
                vals['scout_track_qty'] = True
            product.write(vals)
        if qty is not None:
            current = product.scout_qty_available
            diff = qty - current
            if float_compare(diff, 0.0, precision_digits=3):
                Move = self.env['scout.store.stock.move'].sudo()
                has_moves = bool(Move.search_count([('product_tmpl_id', '=', product.id)]))
                Move.create({
                    'product_tmpl_id': product.id, 'move_kind': 'adjust' if has_moves else 'initial',
                    'quantity': diff, 'unit_cost': product.standard_price, 'date': self.inventory_date,
                    'note': _('Carga desde Excel (%s)', self.file_name or 'xlsx'),
                })
        return created

    def _run(self):
        self.ensure_one()
        rows = self._read_rows()
        if not rows:
            raise UserError(_('El archivo no tiene productos.'))
        log, created, updated, errors, seen = [], 0, 0, 0, {}
        for number, item in rows:
            key = _norm(item['name'])
            if key in seen:
                errors += 1
                log.append(_('Fila %(n)s (%(name)s): repetido (ya venía en la fila %(first)s); se omite.',
                             n=number, name=item['name'], first=seen[key]))
                continue
            seen[key] = number
            try:
                with self.env.cr.savepoint():
                    is_new = self._import_row(item)
                created += is_new
                updated += not is_new
            except Exception as exc:  # noqa: BLE001
                errors += 1
                msg = exc.args[0] if getattr(exc, 'args', None) else str(exc)
                log.append(_('Fila %(n)s (%(name)s): %(msg)s', n=number, name=item['name'], msg=msg))
        header = _('%(c)s nuevos, %(u)s actualizados, %(e)s con error.', c=created, u=updated, e=errors)
        return header + ('\n\n' + '\n'.join(log) if log else _('\n\nSin observaciones.')), created, updated, errors

    def action_check(self):
        self.ensure_one()
        savepoint = self.env.cr.savepoint(flush=True)
        try:
            result, created, updated, errors = self._run()
            self.env.flush_all()
        finally:
            savepoint.close(rollback=True)
            self.env.invalidate_all()
        self.write({'state': 'checked', 'result': _('REVISIÓN (no se guardó nada)\n') + result,
                    'created_count': created, 'updated_count': updated, 'error_count': errors})
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        result, created, updated, errors = self._run()
        self.write({'state': 'done', 'result': result, 'created_count': created,
                    'updated_count': updated, 'error_count': errors})
        return self._reopen()

    def action_view_products(self):
        return self.env['ir.actions.act_window']._for_xml_id('scout_management.action_scout_store_products')

    def _reopen(self):
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}
