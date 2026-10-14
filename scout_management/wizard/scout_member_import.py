# -*- coding: utf-8 -*-
"""Carga masiva de miembros desde Excel (.xlsx)."""
import base64
import io
import logging
import re
import unicodedata
from datetime import date, datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

# (clave interna, encabezado en la plantilla, ancho, ejemplo)
COLUMNS = [
    ('first_name', 'Nombre(s) *', 18, 'Aarón Santiago'),
    ('last_name', 'Apellido paterno', 16, 'Santana'),
    ('last_name2', 'Apellido materno', 16, 'Aparicio'),
    ('scout_name', 'Nombre scout', 14, ''),
    ('birthdate', 'Fecha de nacimiento (dd/mm/aaaa)', 16, '17/05/2021'),
    ('gender', 'Sexo (M/F)', 8, 'M'),
    ('shirt_size', 'Talla', 8, '6'),
    ('type', 'Tipo (Miembro de sección / Dirigente)', 20, 'Miembro de sección'),
    ('section', 'Sección *', 18, 'Castores'),
    ('patrol', 'Equipo (madriguera / seisena / patrulla)', 20, 'Madriguera Roja'),
    ('position', 'Rol', 20, 'Integrante'),
    ('leading_sections', 'Secciones a cargo (dirigentes, separadas por coma)', 22, ''),
    ('state', 'Estado (Activo / En prueba / Baja)', 14, 'Activo'),
    ('join_date', 'Fecha de ingreso', 14, '15/02/2025'),
    ('street', 'Calle y número', 30, '115 oriente 1217 casa 2'),
    ('street2', 'Colonia', 22, 'Bosques de los Héroes sección uno'),
    ('city', 'Ciudad', 14, 'Puebla'),
    ('zip', 'C.P.', 8, '72000'),
    ('phone', 'Teléfono del miembro', 14, ''),
    ('email', 'Correo del miembro', 20, ''),
    ('guardian_name', 'Padre, madre o tutor', 24, 'Luciana Aparicio Santiago'),
    ('guardian_rel', 'Parentesco', 10, 'Madre'),
    ('guardian_phone', 'Teléfono del tutor', 14, '5650010780'),
    ('guardian_phone_alt', 'Teléfono alterno', 14, '2224563489'),
    ('guardian_email', 'Correo del tutor', 22, 'luapari@gmail.com'),
    ('guardian_same_address', 'Tutor vive en el mismo domicilio (Sí/No)', 12, 'Sí'),
    ('guardian_address', 'Domicilio del tutor (si es distinto)', 26, ''),
    ('guardian2_name', 'Segundo padre / madre', 22, ''),
    ('guardian2_rel', 'Parentesco (2)', 10, ''),
    ('guardian2_phone', 'Teléfono (2)', 14, ''),
    ('emergency_name', 'Contacto de emergencia', 22, 'Abuela Santiago'),
    ('emergency_rel', 'Parentesco emergencia', 12, 'Abuela'),
    ('emergency_phone', 'Teléfono emergencia', 14, '2224563489'),
    ('medical_service', 'Servicio médico (IMSS / ISSSTE / Particular / No tengo)', 14, 'No tengo'),
    ('medical_policy', 'No. afiliación o póliza', 14, '000000'),
    ('blood_type', 'Tipo de sangre (O+, A-, …)', 10, 'O+'),
    ('allergies', 'Alergias', 20, 'Miel, abejas'),
    ('medications', 'Medicamentos permanentes', 20, ''),
    ('conditions', 'Condiciones médicas', 22, ''),
    ('promise_date', 'Fecha de Promesa', 14, ''),
    ('last_stage', 'Último adelanto', 18, ''),
    ('last_stage_date', 'Fecha del último adelanto', 14, ''),
    ('transition_date', 'Inicio de ruta de enlace', 14, ''),
    ('reg_cycle', 'Ciclo de registro (ej. 2026)', 10, '2026'),
    ('reg_date', 'Fecha de la ficha', 14, '12/10/2026'),
    ('reg_payment', 'Pago de registro (Pagado / Pendiente / Parcial / Exento)', 14, 'Pagado'),
    ('notes', 'Observaciones', 30, ''),
]

BLOOD = {'o+': 'o+', 'o-': 'o-', 'a+': 'a+', 'a-': 'a-', 'b+': 'b+', 'b-': 'b-', 'ab+': 'ab+', 'ab-': 'ab-',
         'orh+': 'o+', 'orh-': 'o-', 'arh+': 'a+', 'arh-': 'a-', 'brh+': 'b+', 'brh-': 'b-', 'abrh+': 'ab+',
         'abrh-': 'ab-'}
SHIRTS = {'4': '4', '6': '6', '8': '8', '10': '10', '12': '12', '14': '14', '16': '16', 'xs': 'xs',
          'ch': 's', 's': 's', 'm': 'm', 'g': 'l', 'l': 'l', 'xg': 'xl', 'xl': 'xl', 'xxg': 'xxl', 'xxl': 'xxl'}


def _norm(value):
    text = unicodedata.normalize('NFKD', str(value or '')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'\s+', ' ', text).strip()


def _yes(value):
    return _norm(value) in ('si', 's', 'x', 'yes', 'true', '1', 'verdadero')


class ScoutMemberImport(models.TransientModel):
    _name = 'scout.member.import'
    _description = 'Carga masiva de miembros'

    file = fields.Binary(string='Archivo Excel (.xlsx)')
    file_name = fields.Char()
    template_file = fields.Binary(string='Plantilla', readonly=True, compute='_compute_template_file')
    template_name = fields.Char(default='plantilla_miembros_scout.xlsx')
    update_existing = fields.Boolean(string='Actualizar miembros existentes', default=True,
                                     help='Si ya existe un miembro con el mismo nombre (y fecha de nacimiento), se actualiza.')
    state = fields.Selection([('draft', 'Cargar'), ('checked', 'Revisado'), ('done', 'Importado')], default='draft')
    result = fields.Text(string='Resultado', readonly=True)
    created_count = fields.Integer(readonly=True)
    updated_count = fields.Integer(readonly=True)
    error_count = fields.Integer(readonly=True)

    # ------------------------------------------------------------------ plantilla
    def _compute_template_file(self):
        content = self._build_template()
        for wizard in self:
            wizard.template_file = base64.b64encode(content)

    @api.model
    def _build_template(self):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.datavalidation import DataValidation

        wb = Workbook()
        ws = wb.active
        ws.title = 'Miembros'
        head_fill = PatternFill('solid', fgColor='1D2B53')
        req_fill = PatternFill('solid', fgColor='8A5A2B')
        for col, (key, label, width, example) in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = Font(bold=True, color='FFFFFF')
            cell.fill = req_fill if label.endswith('*') else head_fill
            cell.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[cell.column_letter].width = width
            ws.cell(row=2, column=col, value=example)
        ws.row_dimensions[1].height = 48
        ws.freeze_panes = 'B2'

        cat = wb.create_sheet('Catálogos')
        catalogs = {
            'Secciones': [s.name for s in self.env['scout.section'].search([])]
                         + [s.unit_name for s in self.env['scout.section'].search([('unit_name', '!=', False)])],
            'Equipos': [f'{p.name} ({p.section_id.name})' for p in self.env['scout.patrol'].search([])],
            'Roles': self.env['scout.position'].search([]).mapped('name'),
            'Adelantos': self.env['scout.progression.stage'].search([]).mapped('name'),
            'Tipo': ['Miembro de sección', 'Dirigente'],
            'Estado': ['Activo', 'En prueba', 'Baja'],
            'Servicio médico': ['IMSS', 'ISSSTE', 'Particular', 'No tengo'],
            'Pago de registro': ['Pagado', 'Pendiente', 'Parcial', 'Exento'],
            'Tipo de sangre': ['O+', 'O-', 'A+', 'A-', 'B+', 'B-', 'AB+', 'AB-'],
        }
        for col, (title, values) in enumerate(catalogs.items(), start=1):
            head = cat.cell(row=1, column=col, value=title)
            head.font = Font(bold=True, color='FFFFFF')
            head.fill = head_fill
            cat.column_dimensions[head.column_letter].width = 28
            for row, value in enumerate(values, start=2):
                cat.cell(row=row, column=col, value=value)

        def add_list(key, cat_col, count):
            col = [c[0] for c in COLUMNS].index(key) + 1
            letter = ws.cell(row=1, column=col).column_letter
            cat_letter = cat.cell(row=1, column=cat_col).column_letter
            dv = DataValidation(type='list', formula1=f"='Catálogos'!${cat_letter}$2:${cat_letter}${max(count, 1) + 1}",
                                allow_blank=True, showErrorMessage=False)
            ws.add_data_validation(dv)
            dv.add(f'{letter}2:{letter}1000')

        keys = list(catalogs)
        add_list('section', keys.index('Secciones') + 1, len(catalogs['Secciones']))
        add_list('position', keys.index('Roles') + 1, len(catalogs['Roles']))
        add_list('type', keys.index('Tipo') + 1, 2)
        add_list('state', keys.index('Estado') + 1, 3)
        add_list('medical_service', keys.index('Servicio médico') + 1, 4)
        add_list('reg_payment', keys.index('Pago de registro') + 1, 4)
        add_list('blood_type', keys.index('Tipo de sangre') + 1, 8)
        add_list('last_stage', keys.index('Adelantos') + 1, len(catalogs['Adelantos']))

        info = wb.create_sheet('Instrucciones')
        for row, text in enumerate([
            'CARGA MASIVA DE MIEMBROS SCOUT',
            '1. Llena una fila por miembro (beneficiario o dirigente) en la hoja "Miembros". La fila 2 es un ejemplo: bórrala o reemplázala.',
            '2. Las columnas en café son obligatorias. Usa los catálogos de la hoja "Catálogos" (listas desplegables).',
            '3. Fechas en formato dd/mm/aaaa. Si el equipo no existe en la sección, se crea.',
            '4. El tutor se busca por correo, teléfono o nombre para no duplicarlo (hermanos comparten tutor).',
            '5. Si llenas "Ciclo de registro", se crea la ficha validada con el pago indicado (sin generar cargo contable).',
            '6. En Odoo: Grupo Scout > Miembros > Carga masiva desde Excel > "Revisar" y luego "Importar".',
        ], start=1):
            info.cell(row=row, column=1, value=text).font = Font(bold=(row == 1))
        info.column_dimensions['A'].width = 130

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    # ------------------------------------------------------------------ lectura
    def _read_rows(self):
        self.ensure_one()
        if not self.file:
            raise UserError(_('Selecciona el archivo de Excel.'))
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(base64.b64decode(self.file)), data_only=True, read_only=True)
        except Exception as exc:  # noqa: BLE001 - mensaje claro al usuario
            raise UserError(_('No se pudo leer el archivo. Usa la plantilla en formato .xlsx. (%s)', exc)) from exc
        ws = wb['Miembros'] if 'Miembros' in wb.sheetnames else wb.worksheets[0]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        header = [_norm(h).rstrip(' *') for h in rows[0]]
        index = {}
        for key, label, _w, _e in COLUMNS:
            norm_label = _norm(label).rstrip(' *')
            short = norm_label.split(' (')[0]
            for i, h in enumerate(header):
                if h == norm_label or h.split(' (')[0] == short:
                    index[key] = i
                    break
        if 'first_name' not in index or 'section' not in index:
            raise UserError(_('El archivo no tiene las columnas "Nombre(s)" y "Sección". Descarga la plantilla.'))
        data = []
        for number, row in enumerate(rows[1:], start=2):
            if not row or all(v in (None, '') for v in row):
                continue
            item = {key: (row[i] if i < len(row) else None) for key, i in index.items()}
            for key, value in item.items():
                if isinstance(value, str):
                    item[key] = value.strip()
                elif isinstance(value, float) and value.is_integer() and key not in ('birthdate',):
                    item[key] = str(int(value))
            data.append((number, item))
        return data

    @staticmethod
    def _date(value, label):
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, date):
            return value
        text = str(value).strip()
        for fmt in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y'):
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                continue
        raise ValueError(_('%(label)s "%(value)s" no es una fecha válida (usa dd/mm/aaaa).', label=label, value=text))

    # ------------------------------------------------------------------ búsquedas
    def _find_section(self, value):
        if not value:
            return self.env['scout.section']
        n = _norm(value)
        for section in self.env['scout.section'].search([]):
            if n in (_norm(section.name), _norm(section.unit_name), _norm(section.code)):
                return section
        for section in self.env['scout.section'].search([]):
            if n in _norm(section.unit_name) or n in _norm(section.name):
                return section
        raise ValueError(_('Sección "%s" no encontrada.', value))

    def _find_position(self, value, scout_type):
        if not value:
            return self.env['scout.position']
        n = _norm(value)
        positions = self.env['scout.position'].search([])
        for pos in positions:
            if _norm(pos.name) == n or _norm(pos.name).split(' (')[0] == n:
                return pos
        keywords = [('subguia', 'subguide'), ('submadera', 'subguide'), ('subseisener', 'subguide'),
                    ('guia', 'guide'), ('madera', 'guide'), ('seisener', 'guide'), ('subjefe de grupo', 'assistant'),
                    ('jefe de grupo', 'group_chief'), ('subjefe', 'assistant'), ('candidato', 'candidate'),
                    ('jefe', 'chief'), ('scouter', 'scouter'), ('secretari', 'secretary'), ('tesorer', 'treasurer'),
                    ('integrante', 'member')]
        for word, kind in keywords:
            if word in n:
                pos = positions.filtered(lambda p: p.kind == kind)[:1]
                if pos:
                    return pos
        raise ValueError(_('Rol "%s" no está en el catálogo (lo agrega el jefe de grupo).', value))

    def _find_stage(self, value):
        if not value:
            return self.env['scout.progression.stage']
        n = _norm(value)
        stages = self.env['scout.progression.stage'].search([])
        stage = stages.filtered(lambda s: _norm(s.name) == n)[:1] or stages.filtered(lambda s: n in _norm(s.name))[:1]
        if not stage:
            raise ValueError(_('Adelanto "%s" no está en el catálogo.', value))
        return stage

    def _find_or_create_partner(self, name, phone=None, email=None, address=None):
        if not name:
            return self.env['res.partner']
        Partner = self.env['res.partner'].sudo()
        partner = Partner
        if email:
            partner = Partner.search([('email', '=ilike', email)], limit=1)
        if not partner and phone:
            partner = Partner.search([('phone', '=', phone)], limit=1)
        if not partner:
            partner = Partner.search([('name', '=ilike', name), ('is_scout', '=', False)], limit=1)
        vals = {k: v for k, v in (('phone', phone), ('email', email)) if v}
        if address:
            vals['street'] = address
        if partner:
            missing = {k: v for k, v in vals.items() if not partner[k]}
            if missing:
                partner.write(missing)
            return partner
        return Partner.create(dict(vals, name=name))

    # ------------------------------------------------------------------ importación
    def _prepare_member_vals(self, item):
        scout_type = 'leader' if 'dirig' in _norm(item.get('type')) else 'youth'
        section = self._find_section(item.get('section'))
        if not section:
            raise ValueError(_('Falta la sección.'))
        vals = {
            'is_scout': True,
            'company_type': 'person',
            'scout_type': scout_type,
            'scout_first_name': item.get('first_name'),
            'scout_last_name': item.get('last_name') or False,
            'scout_last_name2': item.get('last_name2') or False,
            'scout_name': item.get('scout_name') or False,
            'scout_birthdate': self._date(item.get('birthdate'), _('Fecha de nacimiento')),
            'scout_section_id': section.id,
            'street': item.get('street') or False,
            'street2': item.get('street2') or False,
            'city': item.get('city') or False,
            'zip': item.get('zip') or False,
            'phone': item.get('phone') or False,
            'email': item.get('email') or False,
            'scout_notes': item.get('notes') or False,
        }
        gender = _norm(item.get('gender'))
        if gender:
            vals['scout_gender'] = 'f' if gender.startswith('f') else 'm'
        size = _norm(item.get('shirt_size')).replace('talla ', '')
        if size:
            vals['scout_shirt_size'] = SHIRTS.get(size, False)
        state = _norm(item.get('state'))
        if state:
            vals['scout_state'] = 'inactive' if state.startswith('baja') else ('trial' if 'prueba' in state else 'active')
        if item.get('join_date'):
            vals['scout_join_date'] = self._date(item.get('join_date'), _('Fecha de ingreso'))
        if item.get('transition_date'):
            vals['scout_transition_date'] = self._date(item.get('transition_date'), _('Inicio de ruta de enlace'))
        position = self._find_position(item.get('position'), scout_type)
        if position:
            vals['scout_position_id'] = position.id
        if item.get('patrol'):
            patrol_name = re.sub(r'\s*\(.*\)$', '', str(item['patrol'])).strip()
            Patrol = self.env['scout.patrol'].sudo()
            patrol = Patrol.search([('section_id', '=', section.id), ('name', '=ilike', patrol_name)], limit=1) \
                or Patrol.create({'name': patrol_name, 'section_id': section.id})
            vals['scout_patrol_id'] = patrol.id
        if item.get('leading_sections'):
            sections = self.env['scout.section']
            for name in re.split(r'[,;/]', str(item['leading_sections'])):
                if name.strip():
                    sections |= self._find_section(name.strip())
            vals['scout_leading_section_ids'] = [fields.Command.set(sections.ids)]
        # Médicos
        service = _norm(item.get('medical_service'))
        if service:
            vals['scout_medical_service'] = ('imss' if 'imss' in service else 'issste' if 'issste' in service
                                             else 'none' if 'no' in service else 'private')
        blood = _norm(item.get('blood_type')).replace(' ', '')
        if blood:
            vals['scout_blood_type'] = BLOOD.get(blood, 'unknown')
        for key, fname in (('medical_policy', 'scout_medical_policy'), ('allergies', 'scout_allergies'),
                           ('medications', 'scout_medications'), ('conditions', 'scout_medical_conditions')):
            if item.get(key):
                vals[fname] = str(item[key])
        # Tutores
        same = item.get('guardian_same_address')
        vals['scout_guardian_same_address'] = True if same in (None, '') else _yes(same)
        guardian = self._find_or_create_partner(
            item.get('guardian_name'), item.get('guardian_phone'), item.get('guardian_email'),
            None if vals['scout_guardian_same_address'] else item.get('guardian_address'))
        if guardian:
            vals['scout_guardian_id'] = guardian.id
            vals['scout_guardian_relationship'] = item.get('guardian_rel') or False
            vals['scout_guardian_phone_alt'] = item.get('guardian_phone_alt') or False
        guardian2 = self._find_or_create_partner(item.get('guardian2_name'), item.get('guardian2_phone'))
        if guardian2:
            vals['scout_guardian2_id'] = guardian2.id
            vals['scout_guardian2_relationship'] = item.get('guardian2_rel') or False
        return vals

    def _import_row(self, item):
        Partner = self.env['res.partner'].with_context(scout_change_approved=True, tracking_disable=True)
        vals = self._prepare_member_vals(item)
        name = ' '.join(p for p in (vals['scout_first_name'], vals.get('scout_last_name'),
                                    vals.get('scout_last_name2')) if p)
        domain = [('is_scout', '=', True), ('name', '=ilike', name)]
        if vals.get('scout_birthdate'):
            domain.append(('scout_birthdate', 'in', [vals['scout_birthdate'], False]))
        member = Partner.search(domain, limit=1)
        if member and not self.update_existing:
            raise ValueError(_('Ya existe el miembro %s.', name))
        created = not member
        if member:
            member.write(vals)
        else:
            vals['name'] = name
            member = Partner.create(vals)
        # Emergencia
        if item.get('emergency_name') and item.get('emergency_phone'):
            if not member.scout_emergency_ids.filtered(lambda e: e.phone == str(item['emergency_phone'])):
                self.env['scout.emergency.contact'].sudo().create({
                    'partner_id': member.id, 'name': item['emergency_name'],
                    'relationship': item.get('emergency_rel') or False, 'phone': str(item['emergency_phone'])})
        # Promesa y último adelanto (crean el historial)
        promise_date = self._date(item.get('promise_date'), _('Fecha de Promesa'))
        if promise_date:
            member.scout_promise_date = promise_date
        stage = self._find_stage(item.get('last_stage'))
        if stage:
            member.write({'scout_last_progress_stage_id': stage.id,
                          'scout_last_progress_date': self._date(item.get('last_stage_date'),
                                                                 _('Fecha del último adelanto'))
                          or fields.Date.context_today(self)})
        # Ficha de registro histórica
        if item.get('reg_cycle'):
            self._import_registration(member, item)
        return created

    def _import_registration(self, member, item):
        cycle_name = str(item['reg_cycle']).strip()
        Cycle = self.env['scout.cycle'].sudo()
        cycle = Cycle.search([('name', '=', cycle_name)], limit=1)
        if not cycle:
            if not cycle_name.isdigit():
                raise ValueError(_('Ciclo "%s" no existe.', cycle_name))
            year = int(cycle_name)
            cycle = Cycle.create({'name': cycle_name, 'date_start': date(year, 1, 1), 'date_end': date(year, 12, 31)})
        Reg = self.env['scout.registration'].sudo()
        if Reg.search_count([('member_id', '=', member.id), ('cycle_id', '=', cycle.id)]):
            return
        pay = _norm(item.get('reg_payment'))
        manual = ('paid' if pay.startswith('pag') else 'partial' if pay.startswith('parc')
                  else 'exempt' if pay.startswith('exe') else 'not_paid')
        med = member.sudo()
        Reg.create({
            'member_id': member.id, 'cycle_id': cycle.id, 'state': 'validated', 'source': 'import',
            'first_name': member.scout_first_name or member.name, 'last_name': member.scout_last_name,
            'last_name2': member.scout_last_name2, 'birthdate': member.scout_birthdate,
            'section_id': member.scout_section_id.id, 'street': member.street,
            'guardian_id': member.scout_guardian_id.id, 'guardian_name': member.scout_guardian_id.name,
            'guardian_relationship': member.scout_guardian_relationship,
            'guardian_phone': member.scout_guardian_id.phone, 'guardian_email': member.scout_guardian_id.email,
            'guardian_phone_alt': member.scout_guardian_phone_alt,
            'medical_service': med.scout_medical_service, 'medical_policy': med.scout_medical_policy,
            'blood_type': med.scout_blood_type, 'medical_conditions': med.scout_medical_conditions,
            'manual_payment_state': manual,
            'signed_date': self._date(item.get('reg_date'), _('Fecha de la ficha')) or fields.Date.context_today(self),
            'consent_participation': True, 'consent_first_aid': True, 'consent_data': True,
            'consent_risk': True, 'consent_update': True,
            'guardian_signed_by': member.scout_guardian_id.name or member.name,
            'validated_by_id': self.env.user.id,
        })

    def _run(self, commit):
        self.ensure_one()
        rows = self._read_rows()
        if not rows:
            raise UserError(_('El archivo no tiene filas para importar.'))
        log, created, updated, errors = [], 0, 0, 0
        for number, item in rows:
            if not item.get('first_name'):
                log.append(_('Fila %s: sin nombre, se omite.', number))
                continue
            try:
                with self.env.cr.savepoint():
                    is_new = self._import_row(item)
                created += is_new
                updated += not is_new
            except Exception as exc:  # noqa: BLE001 - se reporta por fila
                errors += 1
                msg = exc.args[0] if getattr(exc, 'args', None) else str(exc)
                log.append(_('Fila %(n)s (%(name)s): %(msg)s', n=number, name=item.get('first_name'), msg=msg))
                _logger.info('Importación scout, fila %s: %s', number, msg)
        header = _('%(c)s nuevos, %(u)s actualizados, %(e)s con error.', c=created, u=updated, e=errors)
        result = header + ('\n\n' + '\n'.join(log) if log else _('\n\nSin observaciones.'))
        return result, created, updated, errors

    def action_check(self):
        """Revisa el archivo sin guardar nada (prueba completa y se deshace)."""
        self.ensure_one()
        savepoint = self.env.cr.savepoint(flush=True)
        try:
            result, created, updated, errors = self._run(commit=False)
            self.env.flush_all()
        finally:
            savepoint.close(rollback=True)
            self.env.invalidate_all()
        self.write({'state': 'checked', 'result': _('REVISIÓN (no se guardó nada)\n') + result,
                    'created_count': created, 'updated_count': updated, 'error_count': errors})
        return self._reopen()

    def action_import(self):
        self.ensure_one()
        result, created, updated, errors = self._run(commit=True)
        self.write({'state': 'done', 'result': result, 'created_count': created,
                    'updated_count': updated, 'error_count': errors})
        return self._reopen()

    def action_view_members(self):
        return self.env['ir.actions.act_window']._for_xml_id('scout_management.action_scout_member')

    def _reopen(self):
        return {'type': 'ir.actions.act_window', 'res_model': self._name, 'res_id': self.id,
                'view_mode': 'form', 'target': 'new'}
