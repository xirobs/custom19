# -*- coding: utf-8 -*-
"""Portal de dirigentes: lo que un jefe de sección captura desde el celular, sin entrar al Odoo completo.

Los jefes de sección son usuarios de portal (gratuitos). Cada método valida que la sección, actividad,
miembro o cargo pertenezca a las secciones del dirigente y luego trabaja con sudo, dejando en el historial
quién hizo el movimiento. Secretaría, tesorería y jefatura siguen trabajando en Odoo con acceso completo.
"""
from datetime import date, datetime, time, timedelta

from dateutil.relativedelta import relativedelta

import pytz

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError

from .calendar_event import CATEGORY_HEX, DEFAULT_TZ
from .scout_common import ACTIVITY_TYPES, quarter_bounds
from .scout_progression import PATH_TYPES

ROLE_ORDER = {'guide': 0, 'subguide': 1}
PATROL_COLORS = {'yellow': '#f2c200', 'blue': '#1f6fb5', 'red': '#d7262f', 'green': '#2e9e44',
                 'orange': '#f36f45', 'purple': '#7b3f98', 'brown': '#8a5a2b', 'white': '#d9dde3'}

QUICK_STATUSES = ('present', 'absent', 'camp', 'outing', 'na')
LEADER_ITEM_TYPES = ('status', 'transition', 'previous', 'objective', 'activity', 'need', 'key', 'progress',
                     'rescue', 'next')


class ResUsers(models.Model):
    _inherit = 'res.users'

    def _scout_leader_sections(self):
        """Secciones que atiende el usuario (vacío = no es dirigente)."""
        self.ensure_one()
        Section = self.env['scout.section'].sudo()
        if not self.share and self.has_group('scout_management.group_scout_officer'):
            return Section.search([])
        partner = self.partner_id.sudo()
        if partner.is_scout and partner.scout_state in ('trial', 'active') \
                and partner.scout_position_id.all_sections_attendance:
            return Section.search([])
        sections = Section.search([('leader_ids', 'in', self.id)])
        sections |= partner.scout_leading_section_ids
        return sections.filtered('active')

    def _scout_is_leader(self):
        self.ensure_one()
        return bool(self._scout_leader_sections())


class ResPartnerLeaderPortal(models.Model):
    _inherit = 'res.partner'

    def action_scout_grant_leader_portal(self):
        """Da acceso al portal de dirigentes (usuario de portal, sin costo de licencia) y lo liga a sus secciones."""
        for partner in self:
            if partner.scout_type != 'leader':
                raise UserError(_('%s no está marcado como dirigente.', partner.name))
            if not partner.scout_leading_section_ids and not partner.scout_section_id:
                raise UserError(_('Indica las "Secciones a cargo" de %s.', partner.name))
            if not partner.email:
                raise UserError(_('%s necesita un correo para recibir su acceso.', partner.name))
            # sudo: con varias empresas, el usuario creado desde otra empresa no se "ve" desde la actual.
            user = partner.sudo().with_context(active_test=False).user_ids[:1]
            if not user:
                same_login = self.env['res.users'].sudo().with_context(active_test=False).search(
                    [('login', '=ilike', (partner.email_normalized or partner.email or '').strip())], limit=1)
                if same_login:
                    raise UserError(_(
                        'El correo %(email)s ya lo usa el usuario "%(user)s", que está ligado a otro contacto '
                        '(%(other)s). Corrige el correo de uno de los dos contactos o usa ese contacto.',
                        email=partner.email, user=same_login.name, other=same_login.partner_id.display_name))
                wizard = self.env['portal.wizard'].sudo().with_context(
                    active_model='res.partner', active_ids=partner.ids).create({})
                wizard.user_ids.action_grant_access()
                user = partner.sudo().with_context(active_test=False).user_ids[:1]
            elif not user.active:
                user.sudo().active = True
            user = user.sudo()
            company = partner.company_id or self.env.company
            if company not in user.company_ids:
                # El dirigente queda también en la empresa del grupo (multiempresa)
                user.write({'company_ids': [fields.Command.link(company.id)]})
            if partner._scout_admin_group():
                # Jefe de grupo, secretaría y tesorería: acceso completo a la parte administrativa del sistema.
                partner._scout_apply_admin_access()
            elif not user.share and not user.has_group('scout_management.group_scout_officer'):
                # Jefe de sección que era usuario interno: pasa a portal (solo ve /dirigentes, sin licencia).
                user.sudo().write({'group_ids': [fields.Command.set([self.env.ref('base.group_portal').id])]})
                partner.message_post(body=_('El usuario %s pasó de usuario interno a usuario de portal.', user.login))
            partner._scout_sync_leader_sections()
            partner.message_post(body=_('Acceso al portal de dirigentes otorgado (%s).', user.login))
        # Sin depender del correo: liga para crear contraseña (WhatsApp / QR) o contraseña asignada aquí.
        wizard = self.env['scout.leader.access.wizard'].create({'partner_id': self[:1].id})
        return wizard._open()

    ADMIN_GROUPS = {'group_chief': 'scout_management.group_scout_manager',
                    'treasurer': 'scout_management.group_scout_treasurer',
                    'secretary': 'scout_management.group_scout_officer'}

    def _scout_admin_group(self):
        """Grupo de acceso administrativo que corresponde al cargo (jefe de grupo, tesorería, secretaría)."""
        self.ensure_one()
        xmlid = self.ADMIN_GROUPS.get(self.sudo().scout_position_id.kind)
        return self.env.ref(xmlid, raise_if_not_found=False) if xmlid else self.env['res.groups']

    def _scout_apply_admin_access(self):
        """Usuario interno con el acceso administrativo de su cargo. El jefe de grupo ve y modifica todo
        (secretaría, tesorería, tienda, aprobaciones). Nunca quita accesos que ya tenga."""
        for partner in self.sudo():
            group = partner._scout_admin_group()
            user = partner.with_context(active_test=False).user_ids[:1]
            if not group or not user or not user.active:
                continue
            if user.share:
                internal = self.env.ref('base.group_user')
                user.write({'group_ids': [fields.Command.set([internal.id, group.id])]})
                partner.message_post(body=_('%(user)s ahora es usuario del sistema con acceso de %(group)s.',
                                            user=user.login, group=group.name))
            elif not user.has_group(group.get_external_id()[group.id]):
                user.write({'group_ids': [fields.Command.link(group.id)]})
                partner.message_post(body=_('%(user)s recibió el acceso de %(group)s.', user=user.login,
                                            group=group.name))

    def _scout_sync_leader_sections(self):
        """Mantiene "Dirigentes con acceso" de cada sección igual a las "Secciones a cargo" del contacto."""
        for partner in self.sudo():
            users = partner.with_context(active_test=False).user_ids
            if not users or partner.scout_type != 'leader':
                continue
            wanted = partner.scout_leading_section_ids
            if not wanted and partner.scout_section_id:
                wanted = partner.scout_section_id
            for user in users:
                current = self.env['scout.section'].sudo().search([('leader_ids', 'in', user.id)])
                for section in wanted - current:
                    section.leader_ids = [fields.Command.link(user.id)]
                if user.share:  # a los usuarios internos no se les quita nada en automático
                    for section in current - wanted:
                        section.leader_ids = [fields.Command.unlink(user.id)]

    def write(self, vals):
        res = super().write(vals)
        if 'scout_leading_section_ids' in vals or 'scout_section_id' in vals:
            self.filtered(lambda p: p.is_scout and p.scout_type == 'leader' and p.user_ids)._scout_sync_leader_sections()
        if 'scout_position_id' in vals:
            # Al nombrarlo jefe de grupo, secretaría o tesorería recibe el acceso administrativo
            self.filtered(lambda p: p.is_scout and p.sudo().user_ids)._scout_apply_admin_access()
        return res


class ScoutLeaderAccessWizard(models.TransientModel):
    _name = 'scout.leader.access.wizard'
    _description = 'Acceso al portal de dirigentes'

    partner_id = fields.Many2one('res.partner', string='Dirigente', required=True, readonly=True)
    login = fields.Char(string='Usuario', readonly=True)
    phone = fields.Char(string='WhatsApp')
    link = fields.Char(string='Liga para crear su contraseña', readonly=True,
                       help='Válida 6 días o hasta que entre por primera vez.')
    qr_code = fields.Binary(string='QR', readonly=True)
    portal_url = fields.Char(string='Dirección del portal', readonly=True)
    password = fields.Char(string='O asigna una contraseña')

    @api.model_create_multi
    def create(self, vals_list):
        wizards = super().create(vals_list)
        from .scout_online_payment import scout_qr_png
        import base64
        for wiz in wizards:
            partner = wiz.partner_id.sudo()
            user = partner.with_context(active_test=False).user_ids[:1]
            partner.signup_prepare(signup_type='signup')
            link = partner.with_context(signup_valid=True)._get_signup_url_for_action(url='/dirigentes')[partner.id]
            base = partner.get_base_url().rstrip('/')
            wiz.write({'login': user.login, 'phone': partner.phone or (partner['mobile'] if 'mobile' in partner._fields else False),
                       'link': link, 'qr_code': base64.b64encode(scout_qr_png(link)),
                       'portal_url': f'{base}/dirigentes'})
        return wizards

    def _open(self):
        return {'type': 'ir.actions.act_window', 'name': _('Acceso al portal de dirigentes'), 'res_model': self._name,
                'res_id': self.id, 'view_mode': 'form', 'target': 'new'}

    def action_whatsapp(self):
        self.ensure_one()
        from urllib.parse import quote
        from .scout_online_payment import scout_whatsapp_phone
        phone = scout_whatsapp_phone(self.phone)
        if not phone:
            raise UserError(_('Escribe el número de WhatsApp del dirigente.'))
        text = _('Hola %(name)s, ya tienes acceso al Portal de dirigentes de %(group)s. '
                 'Entra a esta liga para crear tu contraseña: %(link)s\n'
                 'Tu usuario es: %(login)s\nDespués entras en: %(url)s',
                 name=self.partner_id.name, group=self.env.company.name, link=self.link,
                 login=self.login, url=self.portal_url)
        self.partner_id.message_post(body=_('Liga de acceso al portal de dirigentes enviada por WhatsApp a %s.', phone))
        return {'type': 'ir.actions.act_url', 'url': f'https://wa.me/{phone}?text={quote(text)}', 'target': 'new'}

    def action_set_password(self):
        self.ensure_one()
        if not self.password or len(self.password) < 6:
            raise UserError(_('La contraseña debe tener al menos 6 caracteres.'))
        user = self.partner_id.sudo().with_context(active_test=False).user_ids[:1]
        user.write({'password': self.password})
        self.partner_id.message_post(body=_('Contraseña del portal de dirigentes asignada por %s.', self.env.user.name))
        login, url = self.login, self.portal_url
        self.password = False
        return {'type': 'ir.actions.client', 'tag': 'display_notification',
                'params': {'type': 'success', 'sticky': True, 'title': _('Contraseña asignada'),
                           'message': _('Usuario: %(login)s — entra en %(url)s', login=login, url=url),
                           'next': {'type': 'ir.actions.act_window_close'}}}


class ScoutLeaderPortal(models.AbstractModel):
    _name = 'scout.leader.portal'
    _description = 'Portal de dirigentes'

    # ------------------------------------------------------------------ acceso
    @api.model
    def _sections(self):
        return self.env.user._scout_leader_sections()

    @api.model
    def _section(self, section_id):
        section = self.env['scout.section'].sudo().browse(int(section_id or 0)).exists()
        if not section or section not in self._sections():
            raise AccessError(_('No tienes acceso a esa sección.'))
        return section

    @api.model
    def _members(self, section, include_leaders=False):
        domain = [('is_scout', '=', True), ('scout_state', 'in', ('trial', 'active'))]
        Partner = self.env['res.partner'].sudo()
        members = Partner.search(domain + [('scout_section_id', '=', section.id), ('scout_type', '=', 'youth')])
        if include_leaders:
            members |= Partner.search(domain + ['|', ('scout_leading_section_ids', 'in', section.ids),
                                                ('scout_section_id', '=', section.id), ('scout_type', '=', 'leader')])
        return members.sorted(lambda p: (p.scout_type == 'leader', p.scout_patrol_id.sequence or 0,
                                         p.scout_patrol_id.name or '', p.name or ''))

    @api.model
    def _member(self, member_id, section=None):
        member = self.env['res.partner'].sudo().browse(int(member_id or 0)).exists()
        sections = section or self._sections()
        if not member or not member.is_scout or not (
                member.scout_section_id in sections or member.scout_leading_section_ids & sections):
            raise AccessError(_('No tienes acceso a ese miembro.'))
        return member

    @api.model
    def _event(self, event_id):
        event = self.env['calendar.event'].sudo().browse(int(event_id or 0)).exists()
        if not event or not event.is_scout_event:
            raise AccessError(_('La actividad no existe.'))
        sections = self._sections()
        if event.scout_scope == 'section' and not (event.scout_section_ids & sections):
            raise AccessError(_('No tienes acceso a esa actividad.'))
        return event

    @api.model
    def _event_editable(self, event):
        """Un dirigente edita las actividades de su(s) sección(es); las de todo el grupo solo si las creó."""
        sections = self._sections()
        if self.env.user.has_group('scout_management.group_scout_officer') and not self.env.user.share:
            return True
        if event.scout_scope == 'section':
            return bool(event.scout_section_ids) and not (event.scout_section_ids - sections)
        return event.create_uid == self.env.user

    def _note(self, record, body):
        record.sudo().message_post(body=_('%(body)s — %(user)s (portal de dirigentes)', body=body,
                                          user=self.env.user.name),
                                   author_id=self.env.user.partner_id.id)

    # ------------------------------------------------------------------ fechas
    @api.model
    def _tz(self):
        return pytz.timezone(self.env.user.tz or DEFAULT_TZ)

    @api.model
    def _today(self):
        return datetime.now(self._tz()).date()

    @api.model
    def _to_utc(self, day, hour_float):
        hours = int(hour_float)
        minutes = int(round((hour_float - hours) * 60))
        local = self._tz().localize(datetime.combine(day, time(hours, minutes)))
        return local.astimezone(pytz.utc).replace(tzinfo=None)

    @api.model
    def _local(self, dt):
        return pytz.utc.localize(dt).astimezone(self._tz()) if dt else None

    # ------------------------------------------------------------------ actividades
    @api.model
    def _events_domain(self, section):
        return [('is_scout_event', '=', True),
                '|', ('scout_section_ids', 'in', section.ids), ('scout_scope', '=', 'group')]

    @api.model
    def today_event(self, section):
        return self.env['calendar.event'].sudo().search(
            self._events_domain(section) + [('scout_date', '=', self._today())], order='start', limit=1)

    @api.model
    def events(self, section, upcoming=True, limit=20):
        today = self._today()
        domain = self._events_domain(section) + [('scout_date', '>=' if upcoming else '<', today)]
        return self.env['calendar.event'].sudo().search(domain, order='start %s' % ('asc' if upcoming else 'desc'),
                                                        limit=limit)

    @api.model
    def create_today_session(self, section):
        today = self._today()
        existing = self.today_event(section)
        if existing:
            return existing
        event = self.save_event(section, None, {
            'name': _('Sesión %s', section.display_name), 'date': today.isoformat(),
            'time_from': '10:00', 'time_to': '13:00', 'activity_type': 'regular'})
        event.write({'scout_plan_state': 'approved', 'scout_plan_note': _('Sesión del día (sin planear)')})
        return event

    @api.model
    def _parse_time(self, value, default):
        try:
            hh, mm = (value or default).split(':')[:2]
            return int(hh) + int(mm) / 60.0
        except (ValueError, AttributeError):
            hh, mm = default.split(':')
            return int(hh) + int(mm) / 60.0

    @api.model
    def save_event(self, section, event, post):
        """Crea o edita una actividad de la sección con lo que llena el dirigente."""
        if event and not self._event_editable(event):
            raise AccessError(_('Esta actividad la edita la jefatura.'))
        name = (post.get('name') or '').strip()
        if not name:
            raise UserError(_('Escribe el nombre de la actividad.'))
        try:
            day = fields.Date.to_date(post.get('date')) or self._today()
        except ValueError:
            raise UserError(_('La fecha no es válida.'))
        start_h = self._parse_time(post.get('time_from'), '10:00')
        stop_h = self._parse_time(post.get('time_to'), '13:00')
        try:
            day_end = fields.Date.to_date(post.get('date_end')) or day
        except ValueError:
            day_end = day
        if day_end < day:
            day_end = day
        if day_end == day and stop_h <= start_h:
            stop_h = start_h + 1
        activity_type = post.get('activity_type') or 'regular'
        if activity_type not in dict(ACTIVITY_TYPES) or activity_type in ('council', 'group'):
            activity_type = 'regular'
        category = post.get('category')
        if category not in ('outing', 'ordinary'):
            category = 'outing' if activity_type in ('camp', 'outing') else 'ordinary'
        leaders = self._members(section, include_leaders=True).filtered(lambda p: p.scout_type == 'leader')
        responsible_ids = [int(x) for x in (post.get('responsible_ids') or []) if str(x).isdigit()]
        vals = {
            'name': name,
            'start': self._to_utc(day, start_h),
            'stop': self._to_utc(day_end, stop_h),
            'allday': False,
            'scout_activity_type': activity_type,
            'scout_category': category,
            'scout_theme': post.get('theme') or False,
            'scout_objective': post.get('objective') or False,
            'location': post.get('location') or False,
            'scout_recognitions': post.get('recognitions') or False,
            'description': post.get('description') or False,
            'scout_responsible_ids': [fields.Command.set(leaders.filtered(lambda p: p.id in responsible_ids).ids)],
        }
        if 'has_done' in post:
            vals['scout_done'] = bool(post.get('done'))
        if 'result' in post:
            vals['scout_result'] = post.get('result') or False
        Event = self.env['calendar.event'].sudo().with_context(
            no_mail_to_attendees=True, dont_notify=True, mail_create_nosubscribe=True)
        if event:
            key_changed = event.scout_plan_state in ('approved', 'reviewed') and (
                event.start != vals['start'] or event.stop != vals['stop'] or category != event.scout_category
                or activity_type != event.scout_activity_type or (vals['location'] or False) != (event.location or False))
            if key_changed:
                vals['scout_plan_state'] = 'proposed'
            event.with_context(no_mail_to_attendees=True, dont_notify=True).write(vals)
            self._note(event, _('Actividad actualizada; regresa a revisión del consejo') if key_changed
                       else _('Actividad actualizada'))
            return event
        vals['scout_plan_state'] = 'draft'
        vals.update({'is_scout_event': True, 'scout_scope': 'section',
                     'scout_section_ids': [fields.Command.set(section.ids)],
                     'user_id': self.env.uid, 'partner_ids': [fields.Command.clear()],
                     'privacy': 'public', 'show_as': 'busy'})
        event = Event.create(vals)
        self._note(event, _('Actividad creada'))
        return event

    # ------------------------------------------------------------------ asistencia
    @api.model
    def attendance_lines(self, event, section):
        event._scout_prune_attendance(include_closed=True)
        if not event.scout_attendance_closed:
            # Carga la lista (o agrega a quien se integró a la sección después de crearla).
            missing = event._scout_members_to_call() - event.scout_attendance_line_ids.member_id
            if missing or not event.scout_attendance_line_ids:
                event.action_scout_load_members()
        lines = event.scout_attendance_line_ids
        if event.scout_scope == 'group':
            members = self._members(section, include_leaders=True)
            lines = lines.filtered(lambda l: l.member_id in members or l.section_id == section)
        return lines.sorted(lambda l: (l.member_type == 'leader', l.patrol_id.sequence or 0,
                                       l.patrol_id.name or '', l.member_id.name or ''))

    @api.model
    def set_attendance(self, line_id, status):
        line = self.env['scout.attendance.line'].sudo().browse(int(line_id or 0)).exists()
        if not line or status not in QUICK_STATUSES:
            raise UserError(_('Dato no válido.'))
        event = self._event(line.event_id.id)
        sections = self._sections()
        if event.scout_scope == 'group' and not (line.section_id in sections or line.member_id.scout_section_id in sections
                                                 or line.member_id.scout_leading_section_ids & sections):
            raise AccessError(_('Ese miembro no es de tu sección.'))
        if event.scout_attendance_closed:
            raise UserError(_('La lista ya está cerrada.'))
        line.status = status
        member = line.member_id
        member.invalidate_recordset(['scout_trial_attendance', 'scout_needs_registration'])
        alert = ''
        if member.scout_type == 'youth' and member.scout_registration_status not in ('paid', 'exempt') \
                and member.scout_trial_attendance:
            alert = _('%s.ª asistencia sin registro', member.scout_trial_attendance)
        return {'status': status, 'present': event.scout_present_count, 'counted': event.scout_counted_count,
                'pct': round(event.scout_attendance_pct), 'alert': alert,
                'needs_registration': member.scout_needs_registration}

    @api.model
    def all_present(self, event, section):
        if event.scout_attendance_closed:
            raise UserError(_('La lista ya está cerrada.'))
        self.attendance_lines(event, section).filtered(lambda l: l.status == 'absent').write({'status': 'present'})

    @api.model
    def close_attendance(self, event, section, result=None):
        self.attendance_lines(event, section)
        vals = {'scout_attendance_closed': True, 'scout_done': True}
        if result:
            vals['scout_result'] = result
        event.write(vals)
        self._note(event, _('Lista cerrada: %(present)s de %(counted)s (%(pct)s %%)', present=event.scout_present_count,
                            counted=event.scout_counted_count, pct=round(event.scout_attendance_pct)))
        return event.scout_attendance_line_ids.member_id._scout_notify_registration_due()

    @api.model
    def reopen_attendance(self, event):
        event.write({'scout_attendance_closed': False})
        self._note(event, _('Lista reabierta'))

    # ------------------------------------------------------------------ adelantos
    @api.model
    def stages(self, section):
        return self.env['scout.progression.stage'].sudo().search(
            ['|', ('section_ids', '=', False), ('section_ids', 'in', section.ids)])

    @api.model
    def grant_progress(self, section, stage_id, member_ids, state='achieved', date=None):
        stage = self.stages(section).filtered(lambda s: s.id == int(stage_id or 0))
        if not stage:
            raise UserError(_('Elige el adelanto, insignia o especialidad.'))
        if state not in ('achieved', 'in_progress'):
            state = 'achieved'
        day = fields.Date.to_date(date) if date else self._today()
        leader = self.env.user.partner_id
        Progress = self.env['scout.member.progress'].sudo()
        done = self.env['res.partner'].sudo()
        for member_id in member_ids:
            member = self._member(member_id, section)
            progress = Progress.search([('member_id', '=', member.id), ('stage_id', '=', stage.id)], limit=1)
            vals = {'state': state, 'date_achieved': day if state == 'achieved' else False}
            if not progress:
                Progress.create({'member_id': member.id, 'stage_id': stage.id, 'date_start': day,
                                 'leader_id': leader.id if leader.is_scout else False, **vals})
            elif progress.state != 'achieved':
                progress.write(vals)
            else:
                continue
            done |= member
        for member in done:
            self._note(member, _('%(stage)s: %(state)s', stage=stage.name,
                                 state=_('logrado') if state == 'achieved' else _('en curso')))
        return done

    # ------------------------------------------------------------------ planificación
    @api.model
    def quarter_period(self, offset=0):
        start, _end = quarter_bounds(self._today())
        start = start + relativedelta(months=3 * int(offset or 0))
        start, end = quarter_bounds(start)
        return start, end

    @api.model
    def cycle_period(self, section):
        cycle = self.env['scout.cycle'].sudo()._get_current(section.company_id)
        if cycle:
            return cycle.date_start, cycle.date_end, cycle.name
        today = self._today()
        return date(today.year, 1, 1), date(today.year, 12, 31), str(today.year)

    @api.model
    def period_events(self, section, start, end):
        return self.env['calendar.event'].sudo().search(
            self._events_domain(section) + [('scout_date', '>=', start), ('scout_date', '<=', end)], order='start')

    @api.model
    def plan_days(self, section, start, end, weekday=5):
        """[(fecha, actividades)] con cada sábado del periodo y los demás días que ya tienen actividad."""
        events = self.period_events(section, start, end)
        by_day = {}
        for event in events:
            by_day.setdefault(event.scout_date, self.env['calendar.event'].sudo())
            by_day[event.scout_date] |= event
        day = start
        while day.weekday() != weekday:
            day += timedelta(days=1)
        while day <= end:
            by_day.setdefault(day, self.env['calendar.event'].sudo())
            day += timedelta(days=7)
        return sorted(by_day.items())

    @api.model
    def plan_months(self, section, start, end):
        """[(primer día del mes, actividades)] del ciclo."""
        events = self.period_events(section, start, end)
        months = []
        month = start.replace(day=1)
        while month <= end:
            nxt = month + relativedelta(months=1)
            months.append((month, events.filtered(lambda e: month <= e.scout_date < nxt)))
            month = nxt
        return months

    @api.model
    def generate_sessions(self, section, start, end, hour_from=10.0, hour_to=13.0):
        """Crea la sesión de cada sábado (desde hoy) que aún no tenga actividad de la sección."""
        start = max(start, self._today())
        if start > end:
            raise UserError(_('Ese periodo ya pasó.'))
        before = self.period_events(section, start, end)
        self.env['scout.event.generate'].sudo().create({
            'section_ids': [fields.Command.set(section.ids)], 'date_from': start, 'date_to': end,
            'weekday': '5', 'hour_start': hour_from, 'hour_end': hour_to, 'name': _('Sesión'),
            'activity_type': 'regular', 'skip_existing': True,
        }).action_generate()
        created = self.period_events(section, start, end) - before
        created.write({'user_id': self.env.uid, 'scout_plan_state': 'draft',
                       'scout_approved_by_id': False, 'scout_approved_date': False})
        return created

    @api.model
    def program_pdf(self, section, start, end, theme=None):
        wizard = self.env['scout.program.report.wizard'].sudo().create({
            'section_id': section.id, 'date_from': start, 'date_to': end, 'theme': theme or False,
            'include_group_events': True})
        pdf, _type = self.env['ir.actions.report'].sudo()._render_qweb_pdf(
            'scout_management.action_report_program', wizard.ids, data={'wizard_id': wizard.id})
        return pdf

    # ------------------------------------------------------------------ control de asistencia
    @api.model
    def attendance_matrix(self, section, start, end):
        """La misma matriz del Excel / PDF de asistencia: miembros × sesiones, con cantidad y porcentaje."""
        return self.env['report.scout_management.report_attendance_matrix'].sudo()._build_section(
            section, start, end, False)

    @api.model
    def attendance_pdf(self, section, start, end):
        wizard = self.env['scout.attendance.report.wizard'].sudo().create({
            'section_ids': [fields.Command.set(section.ids)], 'date_from': start, 'date_to': end})
        pdf, _type = self.env['ir.actions.report'].sudo()._render_qweb_pdf(
            'scout_management.action_report_attendance', wizard.ids, data={'wizard_id': wizard.id})
        return pdf

    # ------------------------------------------------------------------ informes trimestral y anual
    @api.model
    def current_report(self, section, create=True, period='quarter', offset=0):
        if period == 'annual':
            start, end, _label = self.cycle_period(section)
        elif period == 'monthly':
            start = self._today().replace(day=1) + relativedelta(months=int(offset or 0))
            end = start + relativedelta(months=1, days=-1)
        else:
            period = 'quarter'
            start, end = self.quarter_period(offset)
        Report = self.env['scout.section.report'].sudo()
        report = Report.search([('section_id', '=', section.id), ('date_from', '=', start),
                                ('period_type', '=', period)], limit=1)
        if not report and create:
            report = Report.create({'section_id': section.id, 'date_from': start, 'date_to': end,
                                    'period_type': period,
                                    'presented_to': _('Consejo de Grupo')})
            report.action_refresh()
        return report

    @api.model
    def save_report(self, report, post):
        if report.state != 'draft':
            raise UserError(_('El informe ya se presentó; pide a secretaría que lo regrese a borrador.'))
        vals = {}
        for key in ('situation', 'closing_phrase'):
            if key in post:
                vals[key] = post.get(key) or False
        commands = []
        for item_type in LEADER_ITEM_TYPES:
            key = f'items_{item_type}'
            if key not in post:
                continue
            for item in report.item_ids.filtered(lambda i: i.item_type == item_type and not i.is_auto):
                commands.append(fields.Command.delete(item.id))
            for seq, text in enumerate(t.strip(' -•\t') for t in (post.get(key) or '').splitlines()):
                if text:
                    commands.append(fields.Command.create({'item_type': item_type, 'name': text,
                                                           'sequence': 100 + seq, 'is_auto': False}))
        if commands:
            vals['item_ids'] = commands
        report.write(vals)
        self._note(report, _('Informe actualizado'))
        return report

    # ------------------------------------------------------------------ mensajes con familias
    @api.model
    def conversations(self, section):
        members = self._members(section)
        Conv = self.env['scout.family.conversation'].sudo()
        convs = Conv.search([('member_id', 'in', members.ids)])
        leader_partners = self.env['res.partner'].sudo().search([('user_ids', 'in', section.leader_ids.ids)])
        leader_partners |= self.env.user.partner_id
        result = []
        for conv in convs:
            last = conv.message_ids.filtered(lambda m: m.message_type == 'comment')[:1]
            if not last:
                continue
            result.append({'conv': conv, 'last': last,
                           'needs_reply': last.author_id not in leader_partners and last.author_id.id
                           in (conv.member_id.scout_guardian_id | conv.member_id.scout_guardian2_id).ids})
        result.sort(key=lambda r: r['last'].date, reverse=True)
        return result

    @api.model
    def conversation(self, conv_id):
        conv = self.env['scout.family.conversation'].sudo().browse(int(conv_id or 0)).exists()
        if not conv:
            raise AccessError(_('La conversación no existe.'))
        self._member(conv.member_id.id)
        conv._sync_followers()
        if self.env.user.partner_id not in conv.message_partner_ids:
            conv.message_subscribe(partner_ids=self.env.user.partner_id.ids)
        conv._portal_ensure_token()
        return conv

    @api.model
    def open_conversation_for_member(self, member_id):
        member = self._member(member_id)
        return self.env['scout.family.conversation'].sudo()._get_or_create_for_member(member)

    # ------------------------------------------------------------------ resumen de la portada
    @api.model
    def dashboard(self, section):
        today_event = self.today_event(section)
        q_start, q_end = self.quarter_period()
        planned = self.period_events(section, q_start, q_end)
        return {
            'today_event': today_event,
            'next_events': self.events(section, upcoming=True, limit=3) - today_event,
            'group_events': self.env['calendar.event'].sudo().search([
                ('is_scout_event', '=', True), ('scout_date', '>=', self._today()),
                ('scout_date', '<=', self._today() + timedelta(days=120)),
                '|', ('scout_scope', '=', 'group'), ('scout_activity_type', 'not in', ('regular', 'none'))],
                order='start', limit=4).filtered(lambda e: e.scout_scope == 'group' or section in e.scout_section_ids
                                                  or section.is_leadership),
            'insurance': self._insurance_installed() and (section.is_leadership or self._is_group_manager()),
            'members': self._members(section),
            'planned_count': len(planned),
            'messages_pending': sum(1 for c in self.conversations(section) if c['needs_reply']),
            'report': self.current_report(section, create=False),
            'annual_report': self.current_report(section, create=False, period='annual'),
        }

    @api.model
    def leader_partners(self, section):
        return self._members(section, include_leaders=True).filtered(lambda p: p.scout_type == 'leader')

    # ------------------------------------------------------------------ equipos (patrulla / seisena / madriguera)
    @api.model
    def team_board(self, section):
        """[(equipo o vacío, miembros)] con guía y subguía primero."""
        members = self._members(section)
        board = []
        for patrol in section.sudo().patrol_ids.sorted(lambda p: (p.sequence, p.id)):
            team = members.filtered(lambda m: m.scout_patrol_id == patrol)
            board.append((patrol, team.sorted(lambda m: (ROLE_ORDER.get(m.scout_position_id.kind, 2), m.name or ''))))
        loose = members.filtered(lambda m: not m.scout_patrol_id or m.scout_patrol_id.section_id != section)
        board.append((self.env['scout.patrol'].sudo(), loose.sorted(lambda m: m.name or '')))
        return board

    @api.model
    def patrol_color(self, patrol):
        if not patrol:
            return '#9aa5b1'
        return patrol.color_hex or PATROL_COLORS.get(patrol.color_name, '#0f3049')

    @api.model
    def edit_patrol(self, section, patrol_id, name, color_hex):
        import re
        patrol = self.env['scout.patrol'].sudo().browse(int(patrol_id or 0)).exists()
        if not patrol or patrol.section_id != section:
            raise AccessError(_('Ese equipo no es de tu sección.'))
        vals = {}
        if (name or '').strip():
            vals['name'] = name.strip()
        if color_hex and re.match(r'^#[0-9a-fA-F]{6}$', color_hex):
            vals['color_hex'] = color_hex.lower()
        if vals:
            patrol.write(vals)
            self._note(section, _('Equipo actualizado: %s', patrol.name))
        return patrol

    @api.model
    def stage_label(self, stage, section):
        """En Tropa la Promesa es la Tercera clase."""
        if stage.stage_type == 'promise' and section.promise_label:
            return section.promise_label
        return stage.name

    @api.model
    def move_member(self, section, member_id, patrol_id):
        member = self._member(member_id, section)
        patrol = self.env['scout.patrol'].sudo()
        if patrol_id:
            patrol = patrol.browse(int(patrol_id)).exists()
            if not patrol or patrol.section_id != section:
                raise AccessError(_('Ese equipo no es de tu sección.'))
        old = member.scout_patrol_id
        member._scout_set_patrol(patrol)
        self._note(member, _('Cambio de %(label)s: %(old)s → %(new)s', label=(section.patrol_label or _('equipo')).lower(),
                             old=old.name or _('sin equipo'), new=patrol.name or _('sin equipo')))
        return member

    @api.model
    def set_role(self, section, member_id, kind):
        member = self._member(member_id, section)
        member._scout_set_patrol_role(kind)
        labels = {'guide': section.patrol_role_label or _('Guía'), 'subguide': section.patrol_subrole_label or _('Subguía'),
                  'member': _('Integrante')}
        self._note(member, _('Cargo en %(team)s: %(role)s', team=member.scout_patrol_id.name or '-', role=labels[kind]))
        return member

    @api.model
    def role_label(self, section, member):
        kind = member.scout_position_id.kind
        if kind == 'guide':
            return section.patrol_role_label or _('Guía')
        if kind == 'subguide':
            return section.patrol_subrole_label or _('Subguía')
        return ''

    @api.model
    def create_patrol(self, section, name, color=None):
        name = (name or '').strip()
        if not name:
            raise UserError(_('Escribe el nombre del equipo.'))
        if color not in PATROL_COLORS:
            color = False
        patrol = self.env['scout.patrol'].sudo().create({
            'name': name, 'section_id': section.id, 'color_name': color,
            'sequence': max(section.sudo().patrol_ids.mapped('sequence') or [0]) + 10})
        self._note(section, _('Nuevo equipo: %s', name))
        return patrol

    # ------------------------------------------------------------------ catálogo de adelantos
    @api.model
    def path_stages(self, section):
        """Promesa, adelantos, ruta de enlace e insignia terminal de la sección, en orden."""
        return self.stages(section).filtered(lambda s: s.stage_type in PATH_TYPES).sorted(
            lambda s: (s.type_order, s.sequence, s.id))

    @api.model
    def specialty_categories(self, section):
        specialties = self.stages(section).filtered(lambda s: s.stage_type == 'specialty')
        categories = specialties.category_id.sorted(lambda c: (c.sequence, c.id))
        result = [(c, specialties.filtered(lambda s: s.category_id == c).sorted(lambda s: (s.sequence, s.id)))
                  for c in categories]
        loose = specialties.filtered(lambda s: not s.category_id)
        if loose:
            result.append((self.env['scout.progression.category'].sudo(), loose))
        return result

    # ------------------------------------------------------------------ especialidades de Tropa
    @api.model
    def specialty_board(self, section):
        """Tablero de especialidades por niveles (Tropa): básicas con amarillo/verde/rojo, técnicas (se piden N)
        y obligatorias de una sola especialidad. None si la sección no tiene especialidades por niveles."""
        Stage = self.env['scout.progression.stage']
        specialties = self.stages(section).filtered(lambda s: s.stage_type == 'specialty')
        leveled = specialties.filtered('level')
        if not leveled:
            return None
        order = {xid: k for k, (xid, _key, _label) in enumerate(Stage.TROPA_BASIC)}
        ref_ids = {}
        for xid, _key, _label in Stage.TROPA_BASIC:
            cat = self.env.ref(f'scout_management.{xid}', raise_if_not_found=False)
            if cat:
                ref_ids[cat.id] = order[xid]
        basic = []
        for cat in leveled.category_id.sorted(lambda c: (ref_ids.get(c.id, 99), c.sequence)):
            levels = {s.level: s for s in leveled.filtered(lambda s, c=cat: s.category_id == c)}
            basic.append({'category': cat, 'levels': [levels.get(lv) for lv in ('1', '2', '3')]})
        tech_cat = self.env.ref('scout_management.stage_cat_tecnicas', raise_if_not_found=False)
        aid_cat = self.env.ref('scout_management.stage_cat_primeros_auxilios', raise_if_not_found=False)
        official = self.env['scout.progression.stage']
        for key, _label in Stage.TROPA_TECHNICAL:
            official |= self.env.ref(f'scout_management.stage_tropa_tec_{key}', raise_if_not_found=False) or Stage
        technical = specialties & official
        technical = technical.sorted(lambda s: [k for k, _l in Stage.TROPA_TECHNICAL].index(
            next((k for k, _l in Stage.TROPA_TECHNICAL if s == self.env.ref(f'scout_management.stage_tropa_tec_{k}',
                                                                            raise_if_not_found=False)), 'campismo')))
        mandatory = specialties.filtered(lambda s: aid_cat and s.category_id == aid_cat)
        path = self.path_stages(section)
        columns = list(path) + [s for b in basic for s in b['levels'] if s] + list(technical) + list(mandatory)
        rows = self.achievements(section, columns)
        need = (tech_cat.required_count or 2) if tech_cat else 2
        for row in rows:
            cells = dict(zip([c.id for c in columns], row['cells']))
            row['by_stage'] = cells
            row['tech_done'] = len([s for s in technical if cells.get(s.id) and cells[s.id].state == 'achieved'])
            row['basic_levels'] = {}
            for b in basic:
                top = 0
                for lv, st in zip((1, 2, 3), b['levels']):
                    if st and cells.get(st.id) and cells[st.id].state == 'achieved':
                        top = lv
                row['basic_levels'][b['category'].id] = top
        return {'basic': basic, 'technical': technical, 'tech_need': need, 'tech_label': tech_cat.name if tech_cat else '',
                'mandatory': mandatory, 'rows': rows, 'path': path}

    # ------------------------------------------------------------------ cuadro de logros
    @api.model
    def achievements(self, section, columns):
        """Filas por miembro: avance del camino de adelantos (termómetro) y celdas de las columnas pedidas."""
        members = self._members(section)
        path = self.path_stages(section)
        progress = self.env['scout.member.progress'].sudo().search([('member_id', 'in', members.ids)])
        by_member = {}
        for p in progress:
            by_member.setdefault(p.member_id.id, {})[p.stage_id.id] = p
        rows = []
        for member in members:
            mine = by_member.get(member.id, {})
            done = [s for s in path if mine.get(s.id) and mine[s.id].state == 'achieved']
            current = next((s for s in path if mine.get(s.id) and mine[s.id].state == 'in_progress'), None)
            if not current:
                current = next((s for s in path if s.id not in mine), None)
            specs = [p for p in mine.values() if p.stage_id.stage_type == 'specialty' and p.state == 'achieved']
            rows.append({
                'member': member, 'cells': [mine.get(c.id) for c in columns],
                'done': len(done), 'total': len(path),
                'pct': round(len(done) * 100.0 / len(path)) if path else 0,
                'next': current, 'specialties': len(specs),
            })
        rows.sort(key=lambda r: (r['member'].scout_patrol_id.sequence or 999, r['member'].scout_patrol_id.name or '',
                                 r['member'].name or ''))
        return rows

    @api.model
    def toggle_progress(self, section, member_id, stage_id):
        """Un toque: vacío → en curso → logrado → vacío."""
        member = self._member(member_id, section)
        stage = self.stages(section).filtered(lambda s: s.id == int(stage_id or 0))
        if not stage:
            raise UserError(_('Ese adelanto no es de tu sección.'))
        Progress = self.env['scout.member.progress'].sudo()
        progress = Progress.search([('member_id', '=', member.id), ('stage_id', '=', stage.id)], limit=1)
        today = self._today()
        if not progress:
            leader = self.env.user.partner_id
            Progress.create({'member_id': member.id, 'stage_id': stage.id, 'state': 'in_progress', 'date_start': today,
                             'leader_id': leader.id if leader.is_scout else False})
            state = 'in_progress'
            self._note(member, _('%s: en curso', stage.name))
        elif progress.state == 'in_progress':
            progress.write({'state': 'achieved', 'date_achieved': today})
            state = 'achieved'
            self._note(member, _('%s: logrado', stage.name))
        else:
            progress.unlink()
            state = 'none'
            self._note(member, _('%s: se quitó el registro', stage.name))
        path = self.path_stages(section)
        achieved = Progress.search_count([('member_id', '=', member.id), ('stage_id', 'in', path.ids),
                                          ('state', '=', 'achieved')])
        return {'state': state, 'done': achieved, 'total': len(path),
                'pct': round(achieved * 100.0 / len(path)) if path else 0}

    # ------------------------------------------------------------------ nuevos ingresos
    @api.model
    def _image_from_post(self, value):
        """Foto tomada con la cámara: llega como data URL (reducida en el celular) o como archivo."""
        import base64
        if hasattr(value, 'read'):
            content = value.read()
            if not content:
                return False
            if len(content) > 8 * 1024 * 1024:
                raise UserError(_('La foto no debe pasar de 8 MB.'))
            return base64.b64encode(content)
        if isinstance(value, str) and value.startswith('data:image'):
            return value.split(',', 1)[1].encode()
        return False

    @api.model
    def _file_from_post(self, value):
        """(base64, nombre) de una foto tomada con la cámara (data URL) o un archivo (imagen o PDF)."""
        import base64
        if hasattr(value, 'read'):
            content = value.read()
            if not content:
                return None
            if len(content) > 8 * 1024 * 1024:
                raise UserError(_('El archivo no debe pasar de 8 MB.'))
            return base64.b64encode(content), value.filename or 'ine'
        if isinstance(value, str) and value.startswith('data:'):
            head, data = value.split(',', 1)
            ext = 'pdf' if 'pdf' in head else 'jpg'
            return data.encode(), f'ine.{ext}'
        return None

    @api.model
    def set_photo(self, section, member_id, value):
        member = self._member(member_id, section)
        image = self._image_from_post(value)
        if not image:
            raise UserError(_('Toma o elige una foto.'))
        member.write({'image_1920': image})
        self._note(member, _('Foto actualizada'))
        return member

    @api.model
    def create_member(self, section, post, photo=None):
        """Alta rápida de un nuevo ingreso (en prueba) para que aparezca en el pase de lista."""
        first = (post.get('first_name') or '').strip()
        last = (post.get('last_name') or '').strip()
        if not first or not last:
            raise UserError(_('Escribe el nombre y el apellido paterno.'))
        try:
            birthdate = fields.Date.to_date(post.get('birthdate')) if post.get('birthdate') else False
        except ValueError:
            raise UserError(_('La fecha de nacimiento no es válida.'))
        Partner = self.env['res.partner'].sudo()
        existing = Partner._scout_find_existing(first, last, post.get('last_name2'), birthdate)
        if existing:
            if existing.scout_state == 'inactive':
                raise UserError(_('%(name)s ya está en el sistema (dado de baja de %(sec)s). Usa "Reingresos" en el '
                                  'bloque de Reinscripción para no duplicarlo.', name=existing.name,
                                  sec=existing.scout_section_id.display_name or '-'))
            if existing.scout_section_id != section:
                raise UserError(_('%(name)s ya está registrado en %(sec)s. Para que pase a tu sección pide a secretaría '
                                  'el cambio de sección; así no se duplica.', name=existing.name,
                                  sec=existing.scout_section_id.display_name or '-'))
            raise UserError(_('%s ya es miembro de tu sección. Búscalo en el bloque de Reinscripción para llenar su '
                              'ficha del nuevo ciclo.', existing.name))
        guardian = Partner
        g_name = (post.get('guardian_name') or '').strip()
        g_phone = (post.get('guardian_phone') or '').strip()
        g_email = (post.get('guardian_email') or '').strip()
        if g_name or g_phone:
            if g_email:
                guardian = Partner.search([('email', '=ilike', g_email)], limit=1)
            if not guardian and g_phone:
                guardian = Partner.search([('phone', '=', g_phone)], limit=1)
            if not guardian:
                guardian = Partner.create({'name': g_name or _('Tutor de %s', first), 'phone': g_phone or False,
                                           'email': g_email or False})
        vals = {
            'is_scout': True, 'scout_type': 'youth', 'scout_state': 'trial',
            'scout_first_name': first, 'scout_last_name': last,
            'scout_last_name2': (post.get('last_name2') or '').strip() or False,
            'name': ' '.join(p for p in (first, last, post.get('last_name2') or '') if p.strip()),
            'scout_name': (post.get('scout_name') or '').strip() or False,
            'scout_birthdate': birthdate, 'scout_gender': post.get('gender') if post.get('gender') in ('m', 'f') else False,
            'scout_section_id': section.id, 'scout_join_date': self._today(),
            'scout_guardian_id': guardian.id or False,
            'scout_guardian_relationship': (post.get('guardian_relationship') or '').strip() or False,
            'street': (post.get('street') or '').strip() or False,
        }
        if post.get('patrol_id') and str(post['patrol_id']).isdigit():
            patrol = self.env['scout.patrol'].sudo().browse(int(post['patrol_id'])).exists()
            if patrol and patrol.section_id == section:
                vals['scout_patrol_id'] = patrol.id
        image = self._image_from_post(photo) if photo else False
        if image:
            vals['image_1920'] = image
        member = Partner.with_context(scout_change_approved=True).create(vals)
        self._note(member, _('Nuevo ingreso registrado en %s', section.display_name))
        return member

    @api.model
    def register_member(self, section, member, post):
        """Ficha de registro firmada por el tutor (consentimientos) y cargo de la cuota con su liga de pago."""
        member = self._member(member.id, section)
        signature = post.get('guardian_signature') or ''
        if not signature.startswith('data:image'):
            raise UserError(_('Falta la firma del padre, madre o tutor.'))
        consents = ('consent_participation', 'consent_first_aid', 'consent_image', 'consent_data', 'consent_risk',
                    'consent_update')
        if not all(post.get(c) for c in consents):
            raise UserError(_('El tutor debe aceptar todas las autorizaciones (incluido el uso de imagen).'))
        reg_existing = self.registration(member)
        ine_front = self._file_from_post(post.get('ine_front_data') or post.get('ine_front'))
        ine_back = self._file_from_post(post.get('ine_back_data') or post.get('ine_back'))
        if not ine_front and not (reg_existing and reg_existing.guardian_id_front):
            raise UserError(_('Adjunta la foto de la identificación del tutor (INE, FM3 o pasaporte).'))
        id_type, id_check = self.env['scout.registration']._scout_verify_id_uploads(
            post.get('guardian_id_type'), [(ine_front, 'front'), (ine_back, 'back')])
        signed_by = (post.get('guardian_signed_by') or '').strip() or member.scout_guardian_id.name
        if not signed_by:
            raise UserError(_('Escribe el nombre de quien firma.'))
        cycle = self.env['scout.cycle'].sudo()._get_registration(section.company_id)
        if not cycle:
            raise UserError(_('No hay un ciclo scout abierto; pide a secretaría que lo cree.'))
        Registration = self.env['scout.registration'].sudo()
        reg = Registration.search([('member_id', '=', member.id), ('cycle_id', '=', cycle.id),
                                   ('state', '!=', 'cancel')], limit=1)
        guardian = member.scout_guardian_id
        vals = {
            'cycle_id': cycle.id, 'member_id': member.id, 'source': 'portal', 'state': 'submitted',
            'company_id': section.company_id.id, 'section_id': section.id,
            'first_name': member.scout_first_name or member.name, 'last_name': member.scout_last_name,
            'last_name2': member.scout_last_name2, 'birthdate': member.scout_birthdate, 'street': member.street,
            'guardian_id': guardian.id or False, 'guardian_name': guardian.name or signed_by,
            'guardian_relationship': member.scout_guardian_relationship, 'guardian_phone': guardian.phone,
            'guardian_email': guardian.email,
            'consent_participation': True, 'consent_first_aid': True, 'consent_data': True,
            'consent_risk': True, 'consent_update': True, 'consent_image': True,
            'guardian_signed_by': signed_by, 'guardian_signature': signature.split(',', 1)[1],
            'signed_date': self._today(), 'leader_user_id': self.env.uid,
        }
        vals.update(guardian_id_type=id_type)
        if id_check:
            vals['guardian_id_check'] = id_check
        if ine_front:
            vals.update(guardian_id_front=ine_front[0], guardian_id_front_name=ine_front[1])
        if ine_back:
            vals.update(guardian_id_back=ine_back[0], guardian_id_back_name=ine_back[1])
        previous = Registration.search_count([('member_id', '=', member.id), ('cycle_id', '!=', cycle.id)])
        vals['kind'] = 'return' if member.scout_state == 'inactive' else ('renewal' if previous else 'new')
        if reg:
            if reg.state == 'validated':
                raise UserError(_('%s ya tiene su ficha validada en este ciclo.', member.name))
            reg.write(vals)
        else:
            reg = Registration.create(vals)
        if not reg.guardian_id:
            reg.guardian_id = reg._get_or_create_guardian()
            if reg.guardian_id and not member.scout_guardian_id:
                member.with_context(scout_change_approved=True).scout_guardian_id = reg.guardian_id
        reg._create_registration_invoice()
        self._note(reg, _('Ficha capturada en el portal de dirigentes con la firma del tutor'))
        officers = self.env.ref('scout_management.group_scout_officer').sudo().all_user_ids.filtered(
            lambda u: u.active and not u.share and u.id != 1)
        todo = self.env.ref('mail.mail_activity_data_todo', raise_if_not_found=False)
        for user in officers[:2]:
            reg.activity_schedule(activity_type_id=todo.id if todo else False,
                                  summary=_('Revisar ficha de nuevo ingreso'), user_id=user.id)
        return reg

    @api.model
    def registration(self, member, cycle=None):
        """Ficha del miembro en el ciclo con inscripciones abiertas (o en el ciclo indicado)."""
        cycle = cycle or self.env['scout.cycle'].sudo()._get_registration()
        return self.env['scout.registration'].sudo().search(
            [('member_id', '=', member.id), ('cycle_id', '=', cycle.id), ('state', '!=', 'cancel')], limit=1) \
            if cycle else self.env['scout.registration'].sudo()

    @api.model
    def payment_info(self, member):
        """Cargo pendiente del registro con su liga, QR y mensaje para WhatsApp."""
        from urllib.parse import quote
        from .scout_online_payment import scout_qr_data_uri, scout_whatsapp_phone
        reg = self.registration(member)
        invoice = reg.invoice_id if reg else self.env['account.move']
        company = (reg.company_id if reg else member.company_id) or self.env.company
        transfer = company._scout_transfer_info(member.name, reg.section_id if reg else member.scout_section_id)
        extra = {}
        if transfer:
            phone = scout_whatsapp_phone(member.scout_guardian_id.phone)
            amount = f' por ${invoice.amount_residual:,.2f}' if invoice else ''
            text = _('Hola, estos son los datos para pagar el registro de %(name)s%(amount)s por transferencia:\n'
                     '%(data)s\nEnvía tu comprobante por WhatsApp. ¡Gracias!', name=member.name, amount=amount,
                     data=transfer['text'])
            extra = {'transfer': transfer, 'whatsapp_transfer': f'https://wa.me/{phone}?text={quote(text)}' if phone
                     else f'https://wa.me/?text={quote(text)}'}
        if not invoice or not invoice.scout_can_pay_online:
            return {'registration': reg, 'invoice': invoice, **extra}
        link = invoice._scout_payment_link()
        phone = scout_whatsapp_phone(invoice.partner_id.phone or member.scout_guardian_id.phone)
        text = invoice._scout_payment_message(link)
        return {**extra, 'registration': reg, 'invoice': invoice, 'link': link, 'qr': scout_qr_data_uri(link),
                'whatsapp': f'https://wa.me/{phone}?text={quote(text)}' if phone else
                f'https://wa.me/?text={quote(text)}'}

    @api.model
    def alerts(self, section):
        """Avisos de la sección: nuevos que ya deben registrarse y miembros por edad."""
        members = self._members(section)
        return {
            'registration': members.filtered('scout_needs_registration'),
            'age': members.filtered(lambda m: m.scout_age_alert in ('over', 'soon')),
        }

    # ------------------------------------------------------------------ planificación: aprobación y calendario general
    @api.model
    def submit_plan(self, section, start, end):
        events = self.period_events(section, start, end).filtered(
            lambda e: e.scout_scope == 'section' and e.scout_plan_state in ('draft', 'changes')
            and self._event_editable(e))
        if not events:
            raise UserError(_('No hay actividades en borrador para enviar.'))
        events.action_scout_plan_submit()
        self._note(section, _('Programa %(start)s – %(end)s enviado al consejo (%(n)s actividades)',
                              start=start.strftime('%d/%m/%Y'), end=end.strftime('%d/%m/%Y'), n=len(events)))
        return events

    @api.model
    def plan_summary(self, events):
        events = events.filtered(lambda e: e.scout_scope == 'section')
        return {state: len(events.filtered(lambda e: e.scout_plan_state == state))
                for state in ('draft', 'proposed', 'reviewed', 'approved', 'changes')}

    @api.model
    def general_calendar(self, month_start, ver='todo'):
        """Semanas del mes con las actividades de todas las secciones y del grupo.
        ver: 'todo', 'grupo' (consejos, eventos, campamentos: sin las juntas ordinarias) o 'juntas'."""
        start = month_start.replace(day=1)
        end = start + relativedelta(months=1, days=-1)
        grid_start = start - timedelta(days=start.weekday())
        grid_end = end + timedelta(days=6 - end.weekday())
        domain = [('is_scout_event', '=', True), ('scout_date', '<=', grid_end),
                  '|', ('scout_date_end', '>=', grid_start), ('scout_date', '>=', grid_start)]
        if ver == 'grupo':
            domain += ['|', ('scout_scope', '=', 'group'), ('scout_activity_type', '!=', 'regular')]
        elif ver == 'juntas':
            domain += [('scout_scope', '=', 'section'), ('scout_activity_type', '=', 'regular')]
        events = self.env['calendar.event'].sudo().search(domain, order='start')
        # Primero lo de todo el grupo y lo especial; después las juntas de cada sección
        events = events.sorted(lambda e: (e.scout_scope != 'group', e.scout_activity_type == 'regular',
                                          e.start or datetime.min))
        weeks, day = [], grid_start
        while day <= grid_end:
            week = []
            for _i in range(7):
                todays = events.filtered(lambda e, d=day: e.scout_date <= d <= (e.scout_date_end or e.scout_date))
                week.append((day, todays))
                day += timedelta(days=1)
            weeks.append(week)
        return {'start': start, 'end': end, 'weeks': weeks, 'events': events}

    @api.model
    def category_hex(self, event):
        return CATEGORY_HEX.get(event.scout_category or 'ordinary')


    # ------------------------------------------------------------------ vista de toda la jefatura de grupo
    @api.model
    def _is_group_manager(self):
        """Jefatura de grupo: oficiales del grupo o quien atiende la sección de jefatura."""
        user = self.env.user
        if not user.share and user.has_group('scout_management.group_scout_officer'):
            return True
        return bool(self._sections().filtered('is_leadership'))

    @api.model
    def _group_sections(self):
        """Secciones de beneficiarios que puede ver la jefatura (sin la propia jefatura)."""
        return self._sections().filtered(lambda s: not s.is_leadership).sorted(lambda s: (s.sequence, s.name or ''))

    @api.model
    def _route_allowed(self, section, route):
        """La sección de la ruta, o la jefatura de grupo sobre cualquier sección que atienda."""
        if not route:
            return False
        if route.section_id == section:
            return True
        return bool(section.is_leadership and route.section_id in self._group_sections())

    GROUP_TYPE_LABELS = [('regular', 'Juntas'), ('council', 'Consejos'), ('camp', 'Campamentos'),
                         ('outing', 'Salidas'), ('group', 'Eventos de grupo'), ('special', 'Especiales'),
                         ('training', 'Formación')]

    @api.model
    def _leader_home_section(self, partner):
        """Sección que atiende un dirigente (la primera que dirige, o la suya)."""
        partner = partner.sudo()
        leading = partner.scout_leading_section_ids.filtered(lambda s: not s.is_leadership)
        return (leading.sorted(lambda s: (s.sequence, s.name or ''))[:1] or partner.scout_section_id)

    @api.model
    def group_attendance(self, start, end):
        """Asistencia de todo el grupo en el periodo (solo listas cerradas): por tipo de actividad, por sección
        (miembros y dirigentes) y de cada jefe y subjefe."""
        from .scout_common import ATTENDED_STATUSES, COUNTED_STATUSES
        types = [t for t, _l in self.GROUP_TYPE_LABELS]
        lines = self.env['scout.attendance.line'].sudo().search([
            ('event_id.is_scout_event', '=', True), ('event_id.scout_attendance_closed', '=', True),
            ('date', '>=', start), ('date', '<=', end)])

        def blank():
            return {'counted': 0, 'attended': 0, 'events': set()}

        def add(bucket, line):
            if line.status in COUNTED_STATUSES:
                bucket['counted'] += 1
                bucket['attended'] += line.status in ATTENDED_STATUSES
            bucket['events'].add(line.event_id.id)

        def pct(bucket):
            return round(bucket['attended'] * 100.0 / bucket['counted']) if bucket['counted'] else None
        by_type = {t: {'youth': blank(), 'leader': blank(), 'all': blank()} for t in types}
        sections = self._group_sections()
        by_section = {sec.id: {'section': sec, 'youth': blank(), 'leader': blank(),
                               'types': {t: blank() for t in types}} for sec in sections}
        leaders = {}
        for line in lines:
            etype = line.event_id.scout_activity_type if line.event_id.scout_activity_type in by_type else 'special'
            kind = 'leader' if line.member_type == 'leader' else 'youth'
            add(by_type[etype][kind], line)
            add(by_type[etype]['all'], line)
            member = line.member_id
            if kind == 'leader':
                sec = self._leader_home_section(member)
                row = leaders.setdefault(member.id, {'member': member, 'section': sec, 'total': blank(),
                                                     'types': {t: blank() for t in types}})
                add(row['total'], line)
                add(row['types'][etype], line)
            else:
                sec = line.section_id or member.scout_section_id
            if sec and sec.id in by_section:
                add(by_section[sec.id][kind], line)
                add(by_section[sec.id]['types'][etype], line)
        # Jefes y subjefes activos aunque todavía no tengan listas cerradas
        active_leaders = self.env['res.partner'].sudo().search([
            ('is_scout', '=', True), ('scout_type', '=', 'leader'), ('scout_state', 'in', ('trial', 'active'))])
        for member in active_leaders:
            leaders.setdefault(member.id, {'member': member, 'section': self._leader_home_section(member),
                                           'total': blank(), 'types': {t: blank() for t in types}})
        leader_rows = []
        for row in leaders.values():
            row['pct'] = pct(row['total'])
            row['type_pct'] = {t: (row['types'][t]['attended'], row['types'][t]['counted']) for t in types}
            leader_rows.append(row)
        order = {sec.id: k for k, sec in enumerate(sections)}
        leader_rows.sort(key=lambda r: (order.get(r['section'].id, 99) if r['section'] else 100,
                                        r['member'].name or ''))
        section_rows = []
        for data in by_section.values():
            data['youth_pct'], data['leader_pct'] = pct(data['youth']), pct(data['leader'])
            data['type_pct'] = {t: pct(data['types'][t]) for t in types}
            data['events'] = len(data['youth']['events'] | data['leader']['events'])
            section_rows.append(data)
        type_rows = []
        for t, label in self.GROUP_TYPE_LABELS:
            data = by_type[t]
            if not data['all']['events']:
                continue
            type_rows.append({'type': t, 'label': label, 'events': len(data['all']['events']),
                              'youth_pct': pct(data['youth']), 'leader_pct': pct(data['leader']),
                              'pct': pct(data['all'])})
        total = blank()
        for line in lines:
            add(total, line)
        youth_all, leader_all = blank(), blank()
        for data in by_type.values():
            for key, target in (('youth', youth_all), ('leader', leader_all)):
                target['counted'] += data[key]['counted']
                target['attended'] += data[key]['attended']
                target['events'] |= data[key]['events']
        return {'types': type_rows, 'sections': section_rows, 'leaders': leader_rows,
                'type_labels': self.GROUP_TYPE_LABELS, 'events': len(total['events']), 'pct': pct(total),
                'youth_pct': pct(youth_all), 'leader_pct': pct(leader_all)}

    @api.model
    def group_attendance_xlsx(self, start, end, label):
        """El reporte general de asistencia en Excel (resumen, secciones y jefes)."""
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        data = self.group_attendance(start, end)
        wb = Workbook()
        navy = PatternFill('solid', fgColor='0F3049')
        head_font = Font(bold=True, color='FFFFFF')

        def sheet(ws, title, headers, rows, widths):
            ws.title = title
            ws.append([f'{self.env.company.name} · Asistencia {label}'])
            ws['A1'].font = Font(bold=True, size=13)
            ws.append(headers)
            for cell in ws[2]:
                cell.fill, cell.font = navy, head_font
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            for row in rows:
                ws.append(row)
            for k, w in enumerate(widths, start=1):
                ws.column_dimensions[ws.cell(row=2, column=k).column_letter].width = w

        def fmt(value):
            return '' if value is None else value / 100.0
        ws = wb.active
        sheet(ws, 'Por tipo', ['Tipo de actividad', 'Actividades', '% miembros', '% dirigentes', '% total'],
              [[r['label'], r['events'], fmt(r['youth_pct']), fmt(r['leader_pct']), fmt(r['pct'])]
               for r in data['types']], [24, 12, 13, 13, 11])
        labels = data['type_labels']
        sheet(wb.create_sheet(), 'Por sección',
              ['Sección', 'Actividades', '% miembros', '% dirigentes'] + [lab for _t, lab in labels],
              [[r['section'].name, r['events'], fmt(r['youth_pct']), fmt(r['leader_pct'])]
               + [fmt(r['type_pct'][t]) for t, _l in labels] for r in data['sections']],
              [24, 12, 13, 13] + [12] * len(labels))
        sheet(wb.create_sheet(), 'Jefes y subjefes',
              ['Dirigente', 'Cargo', 'Sección', 'Asistió', 'De', '%'] + [lab for _t, lab in labels],
              [[r['member'].name, r['member'].scout_position_id.name or '', r['section'].name if r['section'] else '',
                r['total']['attended'], r['total']['counted'], fmt(r['pct'])]
               + ['%s/%s' % r['type_pct'][t] if r['type_pct'][t][1] else '' for t, _l in labels]
               for r in data['leaders']], [32, 20, 20, 9, 7, 8] + [11] * len(labels))
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=3):
                for cell in row:
                    if isinstance(cell.value, float):
                        cell.number_format = '0%'
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    @api.model
    def _insurance_installed(self):
        return bool(self.env['ir.module.module'].sudo().search_count(
            [('name', '=', 'red_seguro_scout'), ('state', '=', 'installed')]))

    LEADERSHIP_KINDS = ('group_chief', 'secretary', 'treasurer')

    @api.model
    def leader_title(self, partner):
        """Dirigentes son solo jefe de grupo, tesorería y secretaría; los demás adultos son scouters."""
        kind = partner.sudo().scout_position_id.kind
        return _('Dirigente') if kind in self.LEADERSHIP_KINDS else _('Scouter')

    @api.model
    def org_chart(self):
        """Organigrama transversal: Consejo de Grupo arriba; al centro la Jefatura de Grupo, a un lado la
        administración (tesorería, secretaría, vocal) y al otro la operación (secciones con sus scouters)."""
        Partner = self.env['res.partner'].sudo()
        leaders = Partner.search([('is_scout', '=', True), ('scout_type', '=', 'leader'),
                                  ('scout_state', 'in', ('trial', 'active'))])
        rank = {'chief': 1, 'assistant': 2, 'scouter': 3, 'candidate': 4}

        def key(p):
            return (rank.get(p.scout_position_id.kind, 9), p.name or '')

        def by_kind(kind):
            return leaders.filtered(lambda p: p.scout_position_id.kind == kind).sorted(key)

        def is_group_level(p):
            home = self._leader_home_section(p)
            return not home or home.is_leadership
        vocal = leaders.filtered(lambda p: 'vocal' in (p.scout_position_id.name or '').lower()).sorted(key)
        sub_group = leaders.filtered(lambda p: p.scout_position_id.kind == 'assistant' and is_group_level(p)
                                     or 'subjefe de grupo' in (p.scout_position_id.name or '').lower()).sorted(key)
        group_people = by_kind('group_chief') | by_kind('treasurer') | by_kind('secretary') | vocal | sub_group
        other_group = leaders.filtered(lambda p: is_group_level(p)) - group_people
        sections = []
        for sec in self._group_sections():
            sec_leaders = leaders.filtered(lambda p, sec=sec: sec in p.scout_leading_section_ids
                                           or p.scout_section_id == sec) - group_people
            youth = Partner.search([('is_scout', '=', True), ('scout_type', '=', 'youth'),
                                    ('scout_state', 'in', ('trial', 'active')), ('scout_section_id', '=', sec.id)])
            patrols = []
            for patrol in youth.scout_patrol_id.sorted(lambda t: (t.sequence, t.name or '')):
                members = youth.filtered(lambda m, t=patrol: m.scout_patrol_id == t)
                patrols.append({'patrol': patrol, 'count': len(members),
                                'guide': members.filtered(lambda m: m.scout_position_id.kind == 'guide')[:1],
                                'subguide': members.filtered(lambda m: m.scout_position_id.kind == 'subguide')[:1]})
            chiefs = sec_leaders.filtered(lambda p: p.scout_position_id.kind == 'chief').sorted(key)
            # Un jefe de grupo que además lleva una sección aparece ahí como interino
            interim = leaders.filtered(lambda p, sec=sec: p in group_people and sec in p.scout_leading_section_ids)
            sections.append({'section': sec, 'chiefs': chiefs, 'interim': interim,
                             'others': (sec_leaders - chiefs).sorted(key), 'youth': len(youth),
                             'patrols': patrols, 'no_patrol': len(youth.filtered(lambda m: not m.scout_patrol_id))})
        return {'chief': by_kind('group_chief'), 'sub_group': sub_group,
                'admin': [('treasurer', _('Tesorería'), 'fa-database', by_kind('treasurer')),
                          ('secretary', _('Secretaría'), 'fa-file-text', by_kind('secretary')),
                          ('vocal', _('Vocal'), 'fa-users', vocal)],
                'other_group': other_group.sorted(key), 'sections': sections, 'leaders': len(leaders)}

    @api.model
    def group_members(self):
        """[(sección, jefes, miembros)] de todas las secciones para la jefatura."""
        result = []
        for sec in self._group_sections():
            members = self._members(sec, include_leaders=True)
            youth = members.filtered(lambda m: m.scout_type == 'youth')
            result.append({'section': sec, 'leaders': members - youth, 'youth': youth,
                           'alerts': len(youth.filtered(lambda m: m.scout_needs_registration or m.scout_age_alert))})
        return result

    @api.model
    def group_resources(self):
        """Por sección: plan de adelantos (recursos ligados a cada adelanto), recursos y videos de actividades."""
        Resource = self.env['scout.resource']
        result = []
        for sec in self._group_sections():
            items = Resource._scout_for_section(sec)
            path = self.path_stages(sec)
            by_stage = []
            for stage in path:
                stage_items = items.filtered(lambda r, st=stage: r.stage_id == st)
                by_stage.append((stage, stage_items))
            staged = items.filtered('stage_id')
            other_staged = staged.filtered(lambda r: r.stage_id not in path)
            videos = (items - staged).filtered(lambda r: r.kind == 'video')
            rest = items - staged - videos
            result.append({'section': sec, 'stages': by_stage, 'other_staged': other_staged, 'videos': videos,
                           'resources': rest, 'total': len(items),
                           'stage_count': len(staged)})
        return result

    @api.model
    def whatsapp_recipients(self, sections, audience='families', text=''):
        """Destinatarios de un aviso por WhatsApp: familias (tutor de cada scout) o jefes de las secciones.
        Cada uno con su liga wa.me y el texto personalizado ({nombre} = nombre del scout o del jefe)."""
        from urllib.parse import quote
        from .scout_online_payment import scout_whatsapp_phone
        Partner = self.env['res.partner'].sudo()
        recipients, seen = [], set()

        def add(person, phone_raw, label, name_for_text, sec):
            phone = scout_whatsapp_phone(phone_raw)
            if not phone or phone in seen:
                if phone in seen:
                    for r in recipients:
                        if r['phone'] == phone and label not in r['label']:
                            r['label'] += ', ' + label
                return
            seen.add(phone)
            body = (text or '').replace('{nombre}', name_for_text or '').replace('{seccion}', sec.name or '')
            recipients.append({'name': person.name, 'label': label, 'phone': phone, 'section': sec,
                               'url': f'https://wa.me/{phone}?text={quote(body)}'})
        missing = []
        for sec in sections:
            if audience == 'leaders':
                leaders = Partner.search([('is_scout', '=', True), ('scout_type', '=', 'leader'),
                                          ('scout_state', 'in', ('trial', 'active')),
                                          '|', ('scout_leading_section_ids', 'in', sec.ids),
                                          ('scout_section_id', '=', sec.id)])
                for leader in leaders.sorted('name'):
                    if not scout_whatsapp_phone(leader.phone):
                        missing.append(leader.name)
                    add(leader, leader.phone, leader.scout_position_id.name or _('Dirigente'),
                        (leader.scout_name or leader.name or '').split()[0] if leader.name else '', sec)
                continue
            for member in self._members(sec).sorted('name'):
                first = (member.scout_name or member.name or '').split()[0] if member.name else ''
                guardian = member.scout_guardian_id
                phone = guardian.phone or member.scout_guardian_phone_alt or (member.phone if not guardian else '')
                if not scout_whatsapp_phone(phone):
                    missing.append(member.name)
                    continue
                add(guardian or member, phone, _('familia de %s', first), first, sec)
        return recipients, missing

    @api.model
    def group_routes(self):
        """[(sección, rutas abiertas ordenadas por urgencia)] de todas las secciones."""
        result = []
        for sec in self._group_sections():
            routes = self.routes(sec)
            result.append((sec, routes))
        return result

    @api.model
    def group_progress(self):
        """Por sección: miembros con su termómetro de adelantos, siguiente paso, especialidades y ruta."""
        result = []
        for sec in self._group_sections():
            rows = self.achievements(sec, [])
            if not rows:
                continue
            routes = {r.member_id.id: r for r in self.env['scout.route'].sudo().search([
                ('member_id', 'in', [r['member'].id for r in rows]), ('is_open', '=', True),
                ('section_id', '=', sec.id)])}
            for row in rows:
                row['route'] = routes.get(row['member'].id)
            rows.sort(key=lambda r: (-r['pct'], r['member'].name or ''))
            total = len(rows)
            result.append({
                'section': sec, 'rows': rows, 'members': total,
                'avg': round(sum(r['pct'] for r in rows) / total) if total else 0,
                'specialties': sum(r['specialties'] for r in rows),
                'no_progress': len([r for r in rows if not r['done']]),
                'red': len([r for r in rows if r['route'] and r['route'].alert == 'red']),
                'yellow': len([r for r in rows if r['route'] and r['route'].alert == 'yellow']),
                'path_total': rows[0]['total'] if rows else 0,
            })
        return result

    # ------------------------------------------------------------------ rutas de enlace y permanencia
    @api.model
    def routes(self, section):
        members = self._members(section)
        members._scout_ensure_route()
        routes = self.env['scout.route'].sudo().search([('member_id', 'in', members.ids), ('is_open', '=', True),
                                                       ('section_id', '=', section.id)])
        return routes.sorted(lambda r: (r.alert_order, r.date_pass_official or date.max, r.member_id.name or ''))

    @api.model
    def incoming_routes(self, section):
        """Beneficiarios que recibirá la sección: en ruta, deben iniciar, vencidos o próximos a comenzar."""
        routes = self.env['scout.route'].sudo().search([
            ('destination_id', '=', section.id), ('is_open', '=', True), ('member_state', 'in', ('trial', 'active')),
            '|', ('state', 'in', ('must_start', 'in_route', 'extension', 'overdue')), ('route_soon', '=', True)])
        return routes.sorted(lambda r: (r.alert_order, r.date_pass_official or date.max))

    @api.model
    def start_route(self, section, route_id, day=None):
        route = self.env['scout.route'].sudo().browse(int(route_id or 0)).exists()
        if not self._route_allowed(section, route):
            raise AccessError(_('Esa ruta no es de tu sección.'))
        self._member(route.member_id.id, route.section_id)
        route.action_start_route(day)
        self._note(route.member_id, _('Inicio de ruta de enlace registrado'))
        return route

    @api.model
    def plan_day(self, section, day):
        """La sesión de la sección en ese día (la crea en borrador si no existe) para planearla y dictar su programa."""
        existing = self.env['calendar.event'].sudo().search([
            ('is_scout_event', '=', True), ('scout_section_ids', 'in', section.ids), ('scout_scope', '=', 'section'),
            ('scout_date', '=', day)], order='start', limit=1)
        if existing:
            return existing
        last = self.env['calendar.event'].sudo().search([
            ('is_scout_event', '=', True), ('scout_section_ids', 'in', section.ids), ('scout_scope', '=', 'section'),
            ('scout_activity_type', '=', 'regular')], order='start desc', limit=1)
        start = self._local(last.start) if last else None
        stop = self._local(last.stop) if last else None
        return self.save_event(section, None, {
            'name': _('Sesión %s', section.display_name), 'date': day.isoformat(),
            'time_from': start.strftime('%H:%M') if start else '16:00',
            'time_to': stop.strftime('%H:%M') if stop and stop.date() == start.date() else '18:00',
            'activity_type': 'regular'})

    # ------------------------------------------------------------------ reingresos y renovaciones
    # ------------------------------------------------------------------ nuevos ingresos y reinscripción
    @api.model
    def enrollment(self, section):
        """Estado de la inscripción del ciclo abierto: reinscripción de activos y nuevos ingresos (en prueba)."""
        cycle = self.env['scout.cycle'].sudo()._get_registration(section.company_id)
        members = self._members(section)
        regs = self.env['scout.registration'].sudo().search(
            [('member_id', 'in', members.ids), ('cycle_id', '=', cycle.id), ('state', '!=', 'cancel')]) \
            if cycle else self.env['scout.registration']
        by_member = {reg.member_id.id: reg for reg in regs}

        def row(member):
            reg = by_member.get(member.id)
            if not reg:
                status = 'none'
            elif reg.payment_state in ('paid', 'exempt'):
                status = 'done'
            else:
                status = 'unpaid'
            return {'member': member, 'reg': reg, 'status': status}

        def group(rows):
            return {key: [r for r in rows if r['status'] == key] for key in ('none', 'unpaid', 'done')}

        renew = group([row(m) for m in members.sorted('name') if m.scout_state == 'active'])
        new = group([row(m) for m in members.sorted('name') if m.scout_state == 'trial'])
        return {
            'cycle': cycle, 'renew': renew, 'new': new,
            'renew_total': sum(len(v) for v in renew.values()), 'new_total': sum(len(v) for v in new.values()),
            'former_count': len(self.former_members(section)),
        }

    @api.model
    def family_form_whatsapp(self, member, cycle):
        """Liga personal de la ficha (sin contraseña) lista para mandar por WhatsApp al tutor."""
        from urllib.parse import quote
        from .scout_online_payment import scout_whatsapp_phone
        member = member.sudo()
        url = member._scout_form_url()
        company = member.company_id or self.env.company
        text = _('Hola, ya está abierta la reinscripción %(cycle)s de %(name)s en %(group)s. Revisa y actualiza su '
                 'ficha aquí (sin contraseña), firma y al final puedes pagar en línea con Mercado Pago (tarjeta u OXXO) '
                 'o en efectivo: %(url)s', cycle=cycle.name if cycle else '', name=member.name,
                 group=company.name, url=url)
        if member.scout_state == 'trial':
            text = _('Hola, esta es la ficha de registro %(cycle)s de %(name)s en %(group)s. Llénala aquí (sin '
                     'contraseña), firma y al final puedes pagar en línea con Mercado Pago (tarjeta u OXXO) o en '
                     'efectivo: %(url)s', cycle=cycle.name if cycle else '', name=member.name, group=company.name,
                     url=url)
        phone = scout_whatsapp_phone(member.scout_guardian_id.phone or member.phone)
        return {'url': url, 'whatsapp': f'https://wa.me/{phone}?text={quote(text)}' if phone
                else f'https://wa.me/?text={quote(text)}'}

    @api.model
    def former_members(self, section):
        return self.env['res.partner'].sudo().search([
            ('is_scout', '=', True), ('scout_type', '=', 'youth'), ('scout_state', '=', 'inactive'),
            ('scout_section_id', '=', section.id)], order='scout_leave_date desc, name')

    @api.model
    def request_return(self, section, member_id, reason):
        member = self._member(member_id, section)
        if member.scout_state != 'inactive':
            raise UserError(_('%s no está dado de baja.', member.name))
        Change = self.env['scout.member.change'].sudo()
        change = Change.search([('member_id', '=', member.id), ('change_type', '=', 'return'),
                                ('state', '=', 'to_approve')], limit=1)
        if not change:
            change = Change.create({
                'member_id': member.id, 'change_type': 'return', 'new_section_id': section.id,
                'reason': (reason or '').strip() or _('Reingreso solicitado desde el portal de dirigentes'),
            })
        self._note(member, _('Reingreso solicitado a %s; falta el visto bueno del jefe de grupo', section.display_name))
        return change

    @api.model
    def registration_label(self, member):
        """Texto del botón de registro: renovar si ya tuvo fichas en ciclos anteriores."""
        current = self.registration(member)
        if current:
            return _('Ficha y firma')
        if member.scout_state == 'inactive':
            return _('Ficha de reingreso')
        cycle = self.env['scout.cycle'].sudo()._get_registration()
        if member.sudo().scout_registration_ids.filtered(lambda r: r.cycle_id != cycle and r.state != 'cancel'):
            return _('Renovar registro %s', cycle.name or '')
        return _('Ficha y firma')

    # ------------------------------------------------------------------ recibir a quien viene de otra sección
    @api.model
    def receive_member(self, section, route_id):
        """El jefe de la sección receptora marca que ya ingresó: pasa a su sección sin patrulla (la asigna después)."""
        route = self.env['scout.route'].sudo().browse(int(route_id or 0)).exists()
        if not route or route.destination_id != section or not route.is_open or route.kind != 'transfer':
            raise AccessError(_('Ese beneficiario no viene a tu sección.'))
        member = route.member_id
        route.with_context(scout_route_receiver=True)._confirm(
            self._today(), section, _('Recibido por %s en el portal de dirigentes.', self.env.user.name))
        self._note(member, _('Ingresó a %s; falta asignarle %s.', section.display_name,
                             (section.patrol_label or _('equipo')).lower()))
        return member

    @api.model
    def pass_options(self, route):
        """Secciones a las que puede pasar (la que corresponde por su sexo va primero)."""
        options = route.section_id._route_destination_options()
        default = route.destination_id or route.section_id._route_destination(route.member_id.scout_gender)
        gender = route.member_id.scout_gender
        if gender:
            options = options.filtered(lambda s: not s.gender or s.gender == gender) or options
        return default, (default | options) if default else options

    @api.model
    def can_give_pass(self, route):
        return route.is_open and route.kind == 'transfer' and route.state in ('must_start', 'in_route', 'extension',
                                                                             'overdue')

    @api.model
    def give_pass(self, section, route_id, destination_id=None, day=None):
        """El jefe de la sección da el pase: el beneficiario pasa a la sección siguiente (sin patrulla)."""
        route = self.env['scout.route'].sudo().browse(int(route_id or 0)).exists()
        if not self._route_allowed(section, route) or route.kind != 'transfer' or not route.is_open:
            raise AccessError(_('Ese beneficiario no está en tu sección.'))
        self._member(route.member_id.id, route.section_id)
        if not self.can_give_pass(route):
            raise UserError(_('%s todavía no tiene edad para el pase de sección.', route.member_id.name))
        _default, options = self.pass_options(route)
        destination = options.filtered(lambda s: s.id == int(destination_id or 0)) if destination_id else _default
        if not destination:
            raise UserError(_('Elige a qué sección pasa.'))
        when = fields.Date.to_date(day) if day else self._today()
        if when > self._today():
            raise UserError(_('La fecha del pase no puede ser futura.'))
        member = route.member_id
        route.with_context(scout_route_receiver=True)._confirm(
            when, destination, _('Pase dado por %s en el portal de dirigentes.', self.env.user.name))
        self._note(member, _('Pase de sección a %s', destination.display_name))
        return member, destination

    @api.model
    def recent_arrivals(self, section, days=90):
        """Recién llegados de otra sección (para resaltarlos y asignarles patrulla)."""
        since = self._today() - timedelta(days=days)
        routes = self.env['scout.route'].sudo().search([('destination_id', '=', section.id), ('is_open', '=', False),
                                                        ('date_pass_real', '>=', since)])
        return routes.member_id.filtered(lambda m: m.scout_section_id == section)

    # ------------------------------------------------------------------ especialidades por colores y técnicas
    LEVEL_COLORS = [('', '#e8edf2', 'Sin nivel'), ('1', '#f2c200', 'Amarillo'), ('2', '#2e9e44', 'Verde'),
                    ('3', '#d7262f', 'Rojo')]

    @api.model
    def specialty_summary(self, section):
        """Por miembro: cuántas especialidades lleva en cada grupo (nivel amarillo / verde / rojo) y sus técnicas."""
        members = self._members(section)
        categories = [c for c, _stages in self.specialty_categories(section) if c]
        progress = self.env['scout.member.progress'].sudo().search(
            [('member_id', 'in', members.ids), ('state', '=', 'achieved'), ('stage_type', '=', 'specialty')])
        rows = []
        for member in members.sorted(lambda m: (m.scout_patrol_id.sequence or 999, m.name or '')):
            mine = progress.filtered(lambda p: p.member_id == member)
            cells = []
            for category in categories:
                done = mine.filtered(lambda p: p.stage_id.category_id == category)
                leveled = done.filtered(lambda p: p.stage_id.level)
                done = done - leveled
                individual = done.filtered(lambda p: not (p.stage_id.name or '').lower().startswith('insignia'))
                level = min(len(individual), 3)
                for p in leveled:
                    level = max(level, int(p.stage_id.level))
                for p in done - individual:
                    name = (p.stage_id.name or '').lower()
                    level = max(level, 3 if 'rojo' in name else 2 if 'verde' in name else 1)
                cells.append({'category': category, 'count': len(individual), 'level': level,
                              'names': ', '.join(individual.stage_id.mapped('name'))})
            rows.append({'member': member, 'cells': cells, 'total': len(mine)})
        return categories, rows

