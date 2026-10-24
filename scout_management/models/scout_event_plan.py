# -*- coding: utf-8 -*-
"""Programa semanal de cada sesión (formato del Grupo 7 Ethnos) y su llenado por dictado."""
import json

from odoo import _, api, fields, models
from odoo.exceptions import AccessError, UserError

from . import scout_plan_parser as parser

FLAG_FIELDS = {
    'structure': 'p_structure', 'symbolic': 'p_symbolic', 'progress': 'p_progress',
    'promise': 'm_promise', 'groups': 'm_groups', 'doing': 'm_doing', 'outdoor': 'm_outdoor', 'adult': 'm_adult',
}


def _hours(value):
    minutes = parser.parse_hhmm(value)
    return minutes / 60.0 if minutes is not None else False


def _hhmm(value):
    if value is False or value is None:
        return ''
    minutes = int(round(value * 60))
    return '%02d:%02d' % divmod(minutes, 60)


class ScoutEventPlanLine(models.Model):
    _name = 'scout.event.plan.line'
    _description = 'Actividad del programa semanal'
    _order = 'event_id, sequence, time_start, id'

    event_id = fields.Many2one('calendar.event', string='Sesión', required=True, ondelete='cascade', index=True)
    sequence = fields.Integer(default=10)
    time_start = fields.Float(string='Inicio')
    time_end = fields.Float(string='Fin')
    duration = fields.Integer(string='Tiempo (min)', compute='_compute_duration', store=True)
    stage = fields.Selection(parser.STAGES, string='Etapa')
    kind = fields.Char(string='Tipo', help='Desfogue, mediano, técnica, reflexión, marco simbólico…')
    name = fields.Char(string='Actividad', required=True)
    objective = fields.Char(string='Objetivo')
    material = fields.Char(string='Material')
    responsible = fields.Char(string='Responsable')
    description = fields.Text(string='Observaciones / descripción')
    p_structure = fields.Boolean(string='Estructura')
    p_symbolic = fields.Boolean(string='Marco simbólico')
    p_progress = fields.Boolean(string='Adelanto progresivo')
    m_promise = fields.Boolean(string='Compromiso (Promesa)')
    m_groups = fields.Boolean(string='Pequeños grupos naturales')
    m_doing = fields.Boolean(string='Aprender haciendo')
    m_outdoor = fields.Boolean(string='Aire libre')
    m_adult = fields.Boolean(string='Adulto estimulante')
    pending = fields.Char(string='Por confirmar')

    @api.depends('time_start', 'time_end')
    def _compute_duration(self):
        for line in self:
            line.duration = int(round((line.time_end - line.time_start) * 60)) \
                if line.time_start and line.time_end and line.time_end > line.time_start else 0

    def _to_dict(self):
        self.ensure_one()
        return parser.new_line(
            id=self.id, start=_hhmm(self.time_start) if self.time_start else '',
            end=_hhmm(self.time_end) if self.time_end else '', stage=self.stage or '', kind=self.kind or '',
            name=self.name or '', objective=self.objective or '', material=self.material or '',
            responsible=self.responsible or '', description=self.description or '',
            flags={key: bool(self[fname]) for key, fname in FLAG_FIELDS.items()},
            pending=[p for p in (self.pending or '').split(' · ') if p])


class CalendarEventPlan(models.Model):
    _inherit = 'calendar.event'

    scout_plan_line_ids = fields.One2many('scout.event.plan.line', 'event_id', string='Programa de la sesión',
                                          copy=True)
    scout_plan_dictation = fields.Text(string='Último dictado', copy=False)
    scout_plan_saved_by_id = fields.Many2one('res.users', string='Programa guardado por', readonly=True, copy=False)

    def _scout_plan_context(self):
        """Hora de inicio de la sesión, lugar y nombres de los dirigentes para el intérprete."""
        self.ensure_one()
        start = self.env['scout.leader.portal']._local(self.start) if self.start else None
        ref = start.hour * 60 + start.minute if start else None
        leaders = self.env['res.partner'].sudo().search([
            ('is_scout', '=', True), ('scout_type', '=', 'leader'),
            '|', ('scout_leading_section_ids', 'in', self.scout_section_ids.ids),
            ('scout_section_id', 'in', self.scout_section_ids.ids)])
        names = []
        for leader in leaders:
            if leader.scout_name:
                names.append(leader.scout_name.strip())
            if leader.scout_first_name:
                names.append(leader.scout_first_name.split()[0])
        return ref, self.location or '', sorted(set(n for n in names if len(n) > 2), key=len, reverse=True)

    # ------------------------------------------------------------------ armar el programa sin capturar todo
    # (nombre, minutos, etapa, tipo, casillas): la junta típica de 2 h 15 min; se ajusta al horario de la sesión.
    BASE_PROGRAM = [
        ('Formación y honores a la bandera', 15, 'opening', 'Ceremonia', ('structure', 'promise')),
        ('Juego de desfogue', 15, 'first', 'Desfogue', ('groups', 'doing', 'outdoor')),
        ('Técnica scout', 30, 'first', 'Técnica', ('progress', 'doing')),
        ('Juego mediano', 20, 'second', 'Mediano', ('groups', 'doing')),
        ('Trabajo por {equipos}: adelantos', 30, 'second', 'Técnica', ('progress', 'groups', 'adult')),
        ('Reflexión y consejo', 15, 'closing', 'Reflexión', ('symbolic', 'adult')),
        ('Cierre, avisos y despedida', 10, 'closing', 'Avisos', ('structure',)),
    ]

    def _scout_plan_window(self):
        """(inicio, fin) de la sesión en minutos del día, hora local."""
        self.ensure_one()
        portal = self.env['scout.leader.portal']
        start = portal._local(self.start) if self.start else None
        stop = portal._local(self.stop) if self.stop else None
        begin = start.hour * 60 + start.minute if start else 15 * 60 + 45
        end = stop.hour * 60 + stop.minute if stop and stop.date() == start.date() else begin + 135
        return begin, max(end, begin + 30)

    def _scout_plan_template(self):
        """Programa base listo para ajustar: las actividades típicas repartidas en el horario de la sesión."""
        self.ensure_one()
        begin, end = self._scout_plan_window()
        total = sum(m for _n, m, *_r in self.BASE_PROGRAM)
        scale = (end - begin) / total
        label = (self.scout_section_ids[:1].patrol_label or 'equipo').lower()
        teams = label + ('es' if label.endswith(('l', 'n', 'r')) else 's')
        lines, cursor = [], begin
        for index, (name, minutes, stage, kind, flags) in enumerate(self.BASE_PROGRAM):
            length = max(5, int(round(minutes * scale / 5.0)) * 5)
            stop = end if index == len(self.BASE_PROGRAM) - 1 else min(cursor + length, end)
            lines.append(parser.new_line(
                start='%02d:%02d' % divmod(cursor, 60), end='%02d:%02d' % divmod(stop, 60), stage=stage, kind=kind,
                name=name.format(equipos=teams), flags={f: f in flags for f in parser.FLAGS}))
            cursor = stop
        return lines

    def _scout_plan_copy_sources(self, limit=12):
        """Otras sesiones de la sección que ya tienen programa (para copiarlo)."""
        self.ensure_one()
        events = self.sudo().search([
            ('id', '!=', self.id), ('is_scout_event', '=', True), ('scout_plan_line_ids', '!=', False),
            ('scout_section_ids', 'in', self.scout_section_ids.ids)], order='start desc', limit=limit)
        portal = self.env['scout.leader.portal']
        return [{'id': e.id, 'label': '%s · %s (%s act.)' % (portal._local(e.start).strftime('%d/%m/%Y'), e.name,
                                                               len(e.scout_plan_line_ids))} for e in events]

    def _scout_plan_copy_from(self, source):
        """Actividades de otra sesión, recorridas al horario de esta."""
        self.ensure_one()
        lines = source.sudo()._scout_plan_lines_dicts()
        begin, _end = self._scout_plan_window()
        first = next((_hours(l['start']) for l in lines if l.get('start')), None)
        offset = begin - int(round(first * 60)) if first is not None else 0

        def shift(value):
            if not value:
                return value
            minutes = int(round(_hours(value) * 60)) + offset
            return '%02d:%02d' % divmod(max(0, min(minutes, 23 * 60 + 59)), 60)

        for line in lines:
            line.pop('id', None)
            line['start'], line['end'] = shift(line.get('start')), shift(line.get('end'))
            line['pending'] = []
        return lines

    def _scout_plan_lines_dicts(self):
        self.ensure_one()
        return [line._to_dict() for line in self.scout_plan_line_ids]

    def scout_plan_interpret(self, text, lines=None):
        self.ensure_one()
        if not (text or '').strip():
            raise UserError(_('Dicta o escribe las actividades.'))
        ref, location, leaders = self._scout_plan_context()
        result = parser.interpret(text, lines if lines is not None else self._scout_plan_lines_dicts(),
                                  ref_start=ref, location=location, leaders=leaders)
        return result

    def scout_plan_save(self, lines, dictation=None, confirm_approved=False):
        """Guarda el programa tal como lo dejó el jefe (reemplaza las actividades anteriores)."""
        self.ensure_one()
        if self.scout_plan_state == 'approved' and self.scout_plan_line_ids and not confirm_approved:
            raise UserError(_('Este programa ya tiene visto bueno. Marca "Confirmo cambiar el programa aprobado" '
                              'para guardar los cambios.'))
        commands = [fields.Command.clear()]
        for seq, line in enumerate(lines or [], start=1):
            name = (line.get('name') or '').strip()
            if not name:
                continue
            vals = {
                'sequence': seq * 10, 'name': name[:200], 'time_start': _hours(line.get('start')),
                'time_end': _hours(line.get('end')),
                'stage': line.get('stage') if line.get('stage') in dict(parser.STAGES) else False,
                'kind': (line.get('kind') or '')[:60] or False, 'objective': (line.get('objective') or '')[:250] or False,
                'material': (line.get('material') or '')[:250] or False,
                'responsible': (line.get('responsible') or '')[:120] or False,
                'description': line.get('description') or False,
                'pending': ' · '.join(p for p in (line.get('pending') or [])
                                      if not (p.startswith('Hora de inicio') and line.get('start'))
                                      and not (p.startswith('Hora de término') and line.get('end'))) or False,
            }
            flags = line.get('flags') or {}
            for key, fname in FLAG_FIELDS.items():
                vals[fname] = bool(flags.get(key))
            commands.append(fields.Command.create(vals))
        vals = {'scout_plan_line_ids': commands, 'scout_plan_saved_by_id': self.env.uid}
        if dictation:
            vals['scout_plan_dictation'] = dictation
        self.with_context(scout_plan_approving=True).write(vals)
        return self.scout_plan_line_ids

    def action_scout_plan_pdf(self):
        return self.env.ref('scout_management.action_report_session_plan').report_action(self)

    def scout_plan_xlsx(self):
        """El programa en el Excel del grupo (mismas 17 columnas que usan las secciones)."""
        self.ensure_one()
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        portal = self.env['scout.leader.portal']
        start = portal._local(self.start) if self.start else None
        wb = Workbook()
        ws = wb.active
        ws.title = (start.strftime('%d-%m-%Y') if start else 'Programa')
        navy, orange = '0F3049', 'F36F45'
        thin = Side(style='thin', color='9AA5B1')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal='center', vertical='center', wrap_text=True)
        wrap = Alignment(vertical='top', wrap_text=True)
        months = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO', 'SEPTIEMBRE', 'OCTUBRE',
                  'NOVIEMBRE', 'DICIEMBRE']
        ws.merge_cells('A1:Q1')
        ws['A1'] = '%s\nPROGRAMA SEMANAL DE ACTIVIDADES' % (self.env.company.name or 'GRUPO SCOUT').upper()
        ws['A1'].font = Font(name='Arial', bold=True, size=14, color='FFFFFF')
        ws['A1'].fill = PatternFill('solid', fgColor=navy)
        ws['A1'].alignment = center
        ws.row_dimensions[1].height = 38
        ws.merge_cells('A2:Q2')
        ws['A2'] = 'SECCIÓN: %s    FECHA: %s' % (
            ', '.join(self.scout_section_ids.mapped('display_name')).upper() or 'GRUPO',
            ('%02d DE %s DE %s %s HRS' % (start.day, months[start.month - 1], start.year, start.strftime('%H:%M')))
            if start else '')
        ws.merge_cells('A3:Q3')
        ws['A3'] = 'ACTIVIDAD: %s\nOBJETIVO: %s\nLUGAR: %s\nENCARGADOS: %s' % (
            (self.name or '').upper(), (self.scout_objective or '').upper(), (self.location or '').upper(),
            ', '.join(self.scout_responsible_ids.mapped(lambda p: p.scout_name or p.name)).upper())
        for ref in ('A2', 'A3'):
            ws[ref].font = Font(name='Arial', bold=True, size=10, color=navy)
            ws[ref].alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[3].height = 60
        ws.merge_cells('J4:L4')
        ws.merge_cells('M4:Q4')
        ws['J4'], ws['M4'] = 'PROGRAMA', 'MÉTODO'
        headers = ['HORA', 'ETAPAS', 'TIPO', 'TIEMPO', 'ACTIVIDAD', 'OBJETIVO', 'MATERIAL', 'RESPONSABLE',
                   'OBSERVACIONES', 'ESTRUCTURA', 'MARCOS SIMBÓLICOS', 'ADELANTO PROGRESIVO', 'COMPROMISO (PROMESA)',
                   'GRUPOS NATURALES PEQUEÑOS', 'MÉTODO APRENDER HACIENDO', 'ACTIVIDAD AL AIRE LIBRE',
                   'ESTIMULANTE Y NO INTERFERENTE DEL ADULTO']
        for col, label in enumerate(headers, start=1):
            cell = ws.cell(row=5, column=col, value=label)
            cell.font = Font(name='Arial', bold=True, size=8, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor=navy if col < 10 else orange)
            cell.alignment = center
            cell.border = border
        for ref in ('J4', 'M4'):
            ws[ref].font = Font(name='Arial', bold=True, color='FFFFFF')
            ws[ref].fill = PatternFill('solid', fgColor=orange)
            ws[ref].alignment = center
        ws.row_dimensions[5].height = 62
        stages = dict(parser.STAGES)
        flag_cols = ['p_structure', 'p_symbolic', 'p_progress', 'm_promise', 'm_groups', 'm_doing', 'm_outdoor',
                     'm_adult']
        row = 6
        for line in self.scout_plan_line_ids:
            hour = _hhmm(line.time_start) if line.time_start else ''
            if line.time_end:
                hour += ' A ' + _hhmm(line.time_end)
            values = [hour or 'POR CONFIRMAR', (stages.get(line.stage) or '').upper(), (line.kind or '').upper(),
                      '%s MIN' % line.duration if line.duration else '', line.name, line.objective or '',
                      line.material or '', line.responsible or '', line.description or '']
            values += ['X' if line[f] else '' for f in flag_cols]
            for col, value in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = Font(name='Arial', size=9, bold=col in (1, 5) or col > 9)
                cell.alignment = center if col in (1, 2, 3, 4) or col > 9 else wrap
                cell.border = border
            row += 1
        widths = [13, 12, 13, 9, 26, 26, 18, 16, 48] + [10] * 8
        for col, width in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=5, column=col).column_letter].width = width
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.fitToWidth = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
