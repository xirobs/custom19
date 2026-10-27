# -*- coding: utf-8 -*-
import base64
import io
from datetime import date

from freezegun import freeze_time

from odoo.exceptions import UserError, ValidationError
from odoo.tests import Form, TransactionCase, tagged


@tagged('post_install', '-at_install', 'red_seguro_scout')
class TestRedSeguroScout(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.policy = cls.env['rss.policy'].create({
            'name': 'Póliza prueba',
            'policy_number': 'T-001',
            'insurer_email': 'aseguradora@example.com',
            'premium_per_member': 300,
            'report_day': 10,
        })
        cls.cluster = cls.env['rss.group.cluster'].create({'name': 'Distrito Prueba'})
        cls.group_a = cls.env['rss.scout.group'].create({
            'name': 'Grupo Prueba A', 'number': 1, 'cluster_id': cls.cluster.id,
            'policy_id': cls.policy.id,
        })
        cls.group_b = cls.env['rss.scout.group'].create({
            'name': 'Grupo Prueba B', 'number': 2, 'cluster_id': cls.cluster.id,
            'policy_id': cls.policy.id,
        })
        cls.Insured = cls.env['rss.insured']

    def _insured(self, group, **kw):
        vals = {
            'first_name': 'Ana', 'last_name_father': 'Pérez', 'last_name_mother': 'López',
            'birthdate': date(2014, 5, 1), 'sex': 'female', 'group_id': group.id,
            'tutor_name': 'Laura López', 'tutor_phone': '2225550000',
        }
        vals.update(kw)
        return self.Insured.create(vals)

    # ------------------------------------------------------------------
    def test_01_name_age_section(self):
        with freeze_time('2026-10-26'):
            ins = self._insured(self.group_a, second_name='  Sofía ',
                                curp='peal140501mplrps09')
        self.assertEqual(ins.name, 'Ana Sofía Pérez López')
        self.assertEqual(ins.curp, 'PEAL140501MPLRPS09')
        self.assertEqual(ins.age, 12)
        self.assertTrue(ins.is_minor)
        self.assertEqual(ins.section, 'tropa')

    def test_02_minor_requires_tutor_to_activate(self):
        ins = self._insured(self.group_a, tutor_name=False, tutor_phone=False)
        with self.assertRaises(ValidationError):
            ins.action_activate()

    def test_03_invalid_email_and_future_birthdate(self):
        with self.assertRaises(ValidationError):
            self._insured(self.group_a, tutor_email='correo-malo')
        with self.assertRaises(ValidationError):
            self._insured(self.group_a, birthdate=date(2999, 1, 1))

    def test_04_group_counts_for_comparison(self):
        self._insured(self.group_a).action_activate()
        self._insured(self.group_a, sex='male', first_name='Luis').action_activate()
        self._insured(self.group_b).action_activate()
        self._insured(self.group_b)  # pre-registro, no cuenta
        self.assertEqual(self.group_a.member_count, 2)
        self.assertEqual(self.group_a.male_count, 1)
        self.assertEqual(self.group_b.member_count, 1)
        self.assertEqual(self.group_b.draft_count, 1)
        self.assertEqual(self.cluster.member_count, 3)

    def test_05_payment_flow_and_receipt(self):
        pending = self._insured(self.group_a)
        pay = self.env['rss.payment'].create({
            'group_id': self.group_a.id, 'amount': 300, 'year': 2026,
            'member_ids': [(6, 0, pending.ids)],
        })
        self.assertEqual(pay.amount_per_member, 300)
        self.assertEqual(pay.amount_expected, 300)
        pay.action_confirm()
        self.assertEqual(pay.state, 'confirmed')
        self.assertTrue(pay.name.startswith('REC/'))
        self.assertEqual(pending.state, 'active', 'Pagar activa los pre-registros')
        self.assertTrue(self.group_a.paid_current_year
                        or date.today().year != 2026)
        with self.assertRaises(UserError):
            pay.action_remit()  # falta folio
        pay.insurer_folio = 'ACUSE-99'
        pay.action_remit()
        self.assertEqual(pay.state, 'remitted')
        with self.assertRaises(UserError):
            pay.action_cancel()
        html = self.env['ir.actions.report']._render_qweb_html(
            'red_seguro_scout.action_report_recibo_pago', pay.ids)[0]
        self.assertIn(b'RECIBO DE PAGO', html)

    def test_06_monthly_report_days_and_movements(self):
        full = self._insured(self.group_a, date_start=date(2026, 1, 1))
        full.state = 'active'
        new = self._insured(self.group_b, first_name='Leo', sex='male',
                            date_start=date(2026, 9, 21))
        new.state = 'active'
        drop = self._insured(self.group_b, first_name='Eva', date_start=date(2026, 1, 1),
                             date_end=date(2026, 9, 10))
        drop.state = 'inactive'
        self._insured(self.group_a, first_name='Fuera', date_start=date(2026, 10, 2)).state = 'active'
        rep = self.env['rss.report'].create({
            'policy_id': self.policy.id, 'period_date': date(2026, 9, 15),
        })
        self.assertEqual(rep.date_from, date(2026, 9, 1))
        self.assertEqual(rep.date_to, date(2026, 9, 30))
        self.assertEqual(rep.deadline, date(2026, 9, 10))
        rep.action_generate()
        self.assertEqual(rep.insured_count, 3)
        days = {l.insured_id: l.days_covered for l in rep.line_ids}
        self.assertEqual(days[full], 30)
        self.assertEqual(days[new], 10)
        self.assertEqual(days[drop], 10)
        self.assertEqual(rep.new_count, 1)
        self.assertEqual(rep.drop_count, 1)
        self.assertEqual(rep.total_days, 50)

        # Excel válido con dos hojas
        rep._generate_xlsx()
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(base64.b64decode(rep.xlsx_file)))
        self.assertEqual(len(book.sheetnames), 2)
        self.assertEqual(book.active.max_row, 5 + 3 + 2)

        # PDF (HTML en modo prueba)
        html = self.env['ir.actions.report']._render_qweb_html(
            'red_seguro_scout.action_report_relacion_mensual', rep.ids)[0]
        self.assertIn(b'Septiembre 2026', html)

        # Envío: abre redactor con Excel + PDF, y al enviar queda "Enviada"
        action = rep.action_send()
        self.assertEqual(len(action['context']['default_attachment_ids']), 2)
        composer = Form(self.env['mail.compose.message'].with_context(action['context']))
        composer.save().action_send_mail()
        self.assertEqual(rep.state, 'sent')

        with self.assertRaises(UserError):
            rep.action_generate()

    def test_07_only_paid_groups(self):
        self._insured(self.group_a, date_start=date(2026, 1, 1)).state = 'active'
        self._insured(self.group_b, date_start=date(2026, 1, 1)).state = 'active'
        self.env['rss.payment'].create({
            'group_id': self.group_a.id, 'amount': 300, 'year': 2026,
            'member_ids': [(6, 0, self.group_a.member_ids.ids)],
        }).action_confirm()
        rep = self.env['rss.report'].create({
            'policy_id': self.policy.id, 'period_date': date(2026, 9, 1),
            'only_paid_groups': True,
        })
        rep.action_generate()
        self.assertEqual(rep.line_ids.group_id, self.group_a)

    def test_08_cron_window_day_1_to_10(self):
        self._insured(self.group_a, date_start=date(2026, 1, 1)).state = 'active'
        Report = self.env['rss.report']
        with freeze_time('2026-11-15'):
            Report._cron_generate_monthly()
        self.assertFalse(Report.search([('policy_id', '=', self.policy.id),
                                        ('auto_generated', '=', True)]))
        with freeze_time('2026-11-03'):
            Report._cron_generate_monthly()
            Report._cron_generate_monthly()
        reps = Report.search([('policy_id', '=', self.policy.id), ('auto_generated', '=', True)])
        self.assertEqual(len(reps), 1, 'No debe duplicar la relación del mes')
        self.assertEqual(reps.period_date, date(2026, 11, 1))
        self.assertEqual(reps.state, 'generated')
        self.assertTrue(reps.activity_ids, 'Deja una actividad para enviarla')

    def test_09_cron_auto_send(self):
        self.policy.auto_send = True
        self._insured(self.group_a, date_start=date(2026, 1, 1)).state = 'active'
        with freeze_time('2026-12-01'):
            self.env['rss.report']._cron_generate_monthly()
        rep = self.env['rss.report'].search([('policy_id', '=', self.policy.id)])
        self.assertEqual(rep.state, 'sent')

    def test_10_pdf_only_attachments_and_wizard(self):
        pdf = b'%PDF-1.4\n%prueba\n'
        wiz = self.env['rss.attach.pdf.wizard'].create({
            'cluster_id': self.cluster.id,
            'group_ids': [(6, 0, (self.group_a | self.group_b).ids)],
            'file': base64.b64encode(pdf), 'filename': 'poliza.pdf',
        })
        wiz.action_attach()
        self.assertEqual(len(self.group_a.attachment_ids), 1)
        self.assertEqual(len(self.group_b.attachment_ids), 1)
        self.assertEqual(len(self.cluster.attachment_ids), 1)
        self.assertEqual(self.group_a.document_count, 2)
        bad = self.env['rss.attach.pdf.wizard'].create({
            'group_ids': [(6, 0, self.group_a.ids)],
            'file': base64.b64encode(b'hola'), 'filename': 'foto.png',
        })
        with self.assertRaises(UserError):
            bad.action_attach()
        png = self.env['ir.attachment'].create({'name': 'foto.png', 'raw': b'\x89PNG',
                                                'mimetype': 'image/png'})
        with self.assertRaises(ValidationError):
            self.group_a.attachment_ids = [(4, png.id)]

    def test_11_claim_coverage(self):
        ins = self._insured(self.group_a, date_start=date(2026, 3, 1))
        ins.action_activate()
        ok = self.env['rss.claim'].create({
            'insured_id': ins.id, 'date_accident': date(2026, 4, 1), 'injury_type': 'fracture',
        })
        late = self.env['rss.claim'].create({
            'insured_id': ins.id, 'date_accident': date(2026, 1, 1), 'injury_type': 'cut',
        })
        self.assertTrue(ok.covered_on_date)
        self.assertFalse(late.covered_on_date)
        self.assertTrue(ok.name.startswith('SIN/'))

    def test_12_group_identity(self):
        """El grupo conserva su personalidad; sin ella cae a la organización y a Red Seguro Scout."""
        from odoo.addons.red_seguro_scout.models.rss_branding import RSS_NAVY, RSS_GOLD
        self.group_a.write({'primary_color': '#1b2a4a', 'secondary_color': '#F07C1E',
                            'motto': 'Siempre listos'})
        with self.assertRaises(ValidationError):
            self.group_b.primary_color = 'azul'
        pay = self.env['rss.payment'].create({
            'group_id': self.group_a.id, 'amount': 300, 'year': 2026,
            'member_ids': [(6, 0, self._insured(self.group_a).ids)],
        })
        brand = pay._rss_brand()
        self.assertEqual(brand['primary'], '#1B2A4A')
        self.assertEqual(brand['on_primary'], '#FFFFFF')
        self.assertEqual(brand['motto'], 'Siempre listos')
        self.assertTrue(brand['own'])
        self.assertIn('--rss-group-primary: #1B2A4A', self.group_a.brand_style)
        self.assertIn('Siempre listos', self.group_a.identity_banner)
        pay.action_confirm()
        html = self.env['ir.actions.report']._render_qweb_html(
            'red_seguro_scout.action_report_recibo_pago', pay.ids)[0]
        self.assertIn(b'#1B2A4A', html)
        self.assertIn(b'Siempre listos', html)
        self.assertIn(b'Red Seguro Scout', html, 'El sello siempre aparece')
        # Grupo sin identidad: organización o, en su defecto, Red Seguro Scout
        self.env.company.primary_color = False
        self.env.company.secondary_color = False
        generic = self.env['rss.branding.mixin']._rss_brand_from(group=self.group_b,
                                                               company=self.env.company)
        self.assertEqual(generic['primary'], RSS_NAVY)
        self.assertEqual(generic['secondary'], RSS_GOLD)
        self.env.company.primary_color = '#7A1F2B'
        self.assertEqual(self.env['rss.branding.mixin']._rss_brand_from(
            group=self.group_b, company=self.env.company)['primary'], '#7A1F2B')

    def test_13_insurance_payment_grants_access(self):
        """Solo quien paga el seguro (a través de su grupo) usa Red Seguro Scout."""
        from odoo.fields import Command
        platform = self.env.ref('red_seguro_scout.bita_platform_red_seguro_scout')
        user = self.env['res.users'].create({
            'name': 'Tesorera Grupo A', 'login': 'tesorera_a@example.com',
            'group_ids': [Command.link(self.env.ref('base.group_user').id)],
        })
        self.group_a.user_ids = [Command.link(user.id)]
        platform.action_enforce()
        rss_user = self.env.ref('red_seguro_scout.group_rss_user')
        with self.assertRaises(ValidationError):
            user.write({'group_ids': [Command.link(rss_user.id)]})
        pay = self.env['rss.payment'].create({
            'group_id': self.group_a.id, 'amount': 300, 'year': date.today().year,
            'member_ids': [(6, 0, self._insured(self.group_a).ids)],
        })
        pay.action_confirm()
        self.assertIn(rss_user, user.group_ids)
        sub = self.env['bita.subscription'].search([('rss_payment_id', '=', pay.id)])
        self.assertEqual(sub.date_end, date(date.today().year, 12, 31))
        self.assertEqual(sub.payment_method, 'group')
        pay.action_cancel()
        self.assertNotIn(rss_user, user.group_ids)
